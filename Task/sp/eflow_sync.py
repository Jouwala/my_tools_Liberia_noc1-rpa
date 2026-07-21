#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
eflow_sync.py
=============
1. Scan the Downloads folder for the latest eFlow export ZIPs whose names
   contain any of the configured prefixes:
       - eflow-pm-task-export-   -> PM task export
       - eflow-cr-export         -> CR (Change Request) export
   The NEWEST file for EACH prefix is selected (so we cover both lists).

2. Each ZIP contains a single .xlsx. Its first sheet's row 1 is the header;
   the remaining rows are the data. The data is loaded into SQLite
   (eflow_pm_cr.sqlite3) into two tables:
       - eflow_pm_tasks
       - eflow_cr_tasks
   The first column (PMT NO. / Change Request Id) is used as the primary key,
   so re-running the script updates existing rows instead of duplicating them.

3. The resulting SQLite database is uploaded to the FTP server.

Dependencies:  pip install openpyxl

Note on credentials: they are hard-coded below for convenience. For better
security, set them via environment variables (EFLOW_FTP_PASS etc.) and read
them with os.environ.get(...).
"""

import os
import io
import re
import sys
import glob
import zipfile
import sqlite3
import datetime
import traceback

import openpyxl

# --------------------------------------------------------------------------- #
# CONFIG  (adjust to taste)
# --------------------------------------------------------------------------- #
DOWNLOAD_DIR = r"C:\Users\ZTEEMS\Downloads"

# The SQLite database is created next to this script by default.
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "eflow_pm_cr.sqlite3")

# Name fragments we look for, mapped to the destination table name.
NAME_PATTERNS = {
    "eflow-pm-task-export-": "eflow_pm_tasks",
    "eflow-cr-export":       "eflow_cr_tasks",
}

# Take the newest N files per pattern? We take 1 per pattern (covers both lists).
FILES_PER_PATTERN = 1

# FTP settings
FTP_HOST = "192.168.100.19"
FTP_PORT = 21
FTP_USER = "ZTEEMS"
FTP_PASS = "DX"                     # consider os.environ["EFLOW_FTP_PASS"]
FTP_REMOTE_DIR = "/01_RPA/LBR_Tools/cr&pm"
FTP_PASSIVE = True

ENABLE_FTP = True                  # set False to only build the DB (dry run)
UPLOAD_DB = True                   # upload eflow_pm_cr.sqlite3
UPLOAD_SOURCE_ZIPS = False         # also upload the original .zip files

# After the DB is written, delete (move to Recycle Bin) every OTHER zip whose
# name contains one of the NAME_PATTERNS, i.e. all matching zips EXCEPT the
# newest one per pattern that was just processed. Files go to the Recycle Bin
# (recoverable) - they are never permanently erased. Set DELETE_OTHER_ZIPS to
# False to disable. With CONFIRM_DELETE True the script lists the files and
# asks for a y/N confirmation before recycling anything.
DELETE_OTHER_ZIPS = True
CONFIRM_DELETE = True
# --------------------------------------------------------------------------- #


def log(msg):
    print(f"[{datetime.datetime.now():%H:%M:%S}] {msg}")


def excel_serial_to_datetime(serial):
    """Convert an Excel serial date number to a datetime."""
    return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=serial)


def normalize_value(v):
    """Turn a raw cell value into a DB-friendly string."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        # Excel serial dates that were not date-formatted show up as floats
        # in the 40000-70000 range. Recover them as dates.
        if 40000 <= v <= 70000:
            try:
                return excel_serial_to_datetime(v).strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                pass
        # Keep integer-looking floats clean
        if v.is_integer():
            return str(int(v))
        return repr(v)
    return str(v)


def sanitize_column(name, used):
    """Make a SQL-safe column name and guarantee uniqueness."""
    col = re.sub(r"[^\w]+", "_", str(name).strip()).strip("_")
    if not col:
        col = "col"
    if col[0].isdigit():
        col = "c_" + col
    base = col
    i = 2
    while col in used:
        col = f"{base}_{i}"
        i += 1
    used.add(col)
    return col


def find_latest_files():
    """Return list of (zip_path, table_name) for the newest file per pattern."""
    selected = []
    for pattern, table in NAME_PATTERNS.items():
        matches = glob.glob(os.path.join(DOWNLOAD_DIR, f"*{pattern}*.zip"))
        if not matches:
            log(f"[WARN] No ZIP matching '{pattern}' found in {DOWNLOAD_DIR}")
            continue
        # sort by modification time, newest first
        matches.sort(key=lambda p: os.path.getmtime(p), reverse=True)
        for path in matches[:FILES_PER_PATTERN]:
            selected.append((path, table))
            log(f"Selected ({table}): {os.path.basename(path)}")
    return selected


def read_xlsx_from_zip(zip_path):
    """Return (headers, rows) from the first .xlsx inside the ZIP."""
    with zipfile.ZipFile(zip_path) as z:
        xlsx_names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        if not xlsx_names:
            raise RuntimeError(f"No .xlsx found inside {zip_path}")
        data = z.read(xlsx_names[0])

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        raise RuntimeError(f"Worksheet is empty in {zip_path}")

    # Detect the header row: skip rows whose first cell looks like a data id
    # (e.g. 'PMT2026...', 'CR2026...'). The header row contains 'NO.'/'Id'/'Status'.
    header_idx = 0
    for i, row in enumerate(all_rows[:5]):
        first = str(row[0]) if row and row[0] is not None else ""
        if re.match(r"^(PMT|CR|PMP|TT)\d", first):
            continue
        header_idx = i
        break

    raw_headers = [str(h).strip() if h is not None else f"col{j}"
                   for j, h in enumerate(all_rows[header_idx])]
    used = set()
    headers = [sanitize_column(h, used) for h in raw_headers]

    data_rows = all_rows[header_idx + 1:]
    data_rows = [r for r in data_rows if any(c is not None and str(c).strip() != "" for c in r)]
    log(f"  parsed {len(headers)} columns, {len(data_rows)} data rows "
        f"(header at sheet row {header_idx + 1})")
    return headers, data_rows


def write_to_db(conn, table, headers, rows, source_file):
    cur = conn.cursor()
    used = set(headers)
    pk = headers[0]
    col_defs = ", ".join(f'"{c}" TEXT' for c in headers)
    cur.execute(
        f'CREATE TABLE IF NOT EXISTS "{table}" ('
        f'{col_defs}, "source_file" TEXT, "imported_at" TEXT, '
        f'PRIMARY KEY ("{pk}"))'
    )

    cols_sql = ", ".join(f'"{c}"' for c in headers)
    placeholders = ", ".join("?" for _ in headers) + ", ?, ?"
    insert_sql = (
        f'INSERT OR REPLACE INTO "{table}" '
        f'({cols_sql}, "source_file", "imported_at") VALUES ({placeholders})'
    )

    imported_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    n = 0
    for r in rows:
        vals = [normalize_value(r[j]) if j < len(r) else "" for j in range(len(headers))]
        cur.execute(insert_sql, vals + [os.path.basename(source_file), imported_at])
        n += 1
    conn.commit()
    log(f"  inserted/replaced {n} rows into '{table}'")


def recycle_file(path):
    """Move a file to the Windows Recycle Bin (safe, recoverable delete).

    Uses the SHFileOperation API so the file can be restored from the Bin.
    Raises on failure - it NEVER falls back to a permanent os.remove().
    """
    if sys.platform != "win32":
        raise RuntimeError("Recycle Bin delete is only supported on Windows")
    import ctypes
    from ctypes import wintypes

    class SHFILEOPSTRUCT(ctypes.Structure):
        _fields_ = [
            ("hwnd", wintypes.HWND),
            ("wFunc", wintypes.UINT),
            ("pFrom", ctypes.c_wchar_p),
            ("pTo", ctypes.c_wchar_p),
            ("fFlags", wintypes.UINT),
            ("fAnyOperationsAborted", wintypes.BOOL),
            ("hNameMappings", ctypes.c_void_p),
            ("lpszProgressTitle", ctypes.c_wchar_p),
        ]

    FO_DELETE = 3
    FOF_SILENT = 4
    FOF_NOCONFIRMATION = 16
    FOF_ALLOWUNDO = 64          # <- send to Recycle Bin instead of permanent delete
    FOF_NOERRORUI = 1024

    # pFrom must be double-null terminated
    p_from = ctypes.c_wchar_p(os.path.abspath(path) + "\0\0")
    op = SHFILEOPSTRUCT(
        0, FO_DELETE, p_from, None,
        FOF_ALLOWUNDO | FOF_NOCONFIRMATION | FOF_SILENT | FOF_NOERRORUI,
        False, None, None,
    )
    res = ctypes.windll.shell32.SHFileOperationW(ctypes.byref(op))
    if res != 0:
        raise RuntimeError(f"SHFileOperation failed (code {res})")


def delete_other_zips(selected_paths):
    """Move every matching zip (except the ones just processed) to Recycle Bin."""
    if not DELETE_OTHER_ZIPS:
        log("DELETE_OTHER_ZIPS is False - skipping deletion of other zips.")
        return

    selected_abs = {os.path.abspath(p) for p in selected_paths}
    to_delete = []
    for pattern in NAME_PATTERNS:
        for p in glob.glob(os.path.join(DOWNLOAD_DIR, f"*{pattern}*.zip")):
            if os.path.abspath(p) not in selected_abs:
                to_delete.append(p)

    if not to_delete:
        log("No other matching zip files to delete.")
        return

    log(f"Found {len(to_delete)} other matching zip file(s) to recycle:")
    for p in to_delete:
        log(f"   - {p}")

    if CONFIRM_DELETE:
        try:
            ans = input("Move these files to the Recycle Bin? [y/N] ").strip().lower()
        except EOFError:
            ans = ""
        if ans not in ("y", "yes"):
            log("Deletion cancelled by user - no files removed.")
            return

    for p in to_delete:
        try:
            recycle_file(p)
            log(f"  recycled: {os.path.basename(p)}")
        except Exception as e:
            log(f"  [ERROR] could not recycle {p}: {e}")


def ensure_remote_dir(ftp, remote_dir):
    """Create the remote directory (and parents) if needed, then cwd into it."""
    parts = [p for p in remote_dir.strip("/").split("/") if p]
    path = ""
    for p in parts:
        path = path + "/" + p
        try:
            ftp.cwd(path)
        except Exception:
            try:
                ftp.mkd(path)
                ftp.cwd(path)
            except Exception:
                pass  # best effort


def upload_to_ftp(files_to_upload):
    from ftplib import FTP
    log(f"Connecting to FTP {FTP_HOST}:{FTP_PORT} ...")
    ftp = FTP()
    ftp.connect(FTP_HOST, FTP_PORT, timeout=30)
    ftp.login(FTP_USER, FTP_PASS)
    ftp.set_pasv(FTP_PASSIVE)
    ensure_remote_dir(ftp, FTP_REMOTE_DIR)
    log(f"Remote working dir: {ftp.pwd()}")

    for local_path in files_to_upload:
        remote_name = os.path.basename(local_path)
        log(f"Uploading {remote_name} ({os.path.getsize(local_path)} bytes) ...")
        with open(local_path, "rb") as f:
            ftp.storbinary(f"STOR {remote_name}", f)
        log(f"  -> uploaded {remote_name}")
    ftp.quit()
    log("FTP upload finished.")


def main():
    selected = find_latest_files()
    if not selected:
        log("Nothing to do - no matching ZIP files found.")
        return

    conn = sqlite3.connect(DB_PATH)
    try:
        upload_list = []
        selected_paths = []
        for zip_path, table in selected:
            log(f"Processing {os.path.basename(zip_path)} -> table '{table}'")
            headers, rows = read_xlsx_from_zip(zip_path)
            write_to_db(conn, table, headers, rows, zip_path)
            selected_paths.append(zip_path)
            if UPLOAD_SOURCE_ZIPS:
                upload_list.append(zip_path)

        # After the DB is written, recycle the other matching zip files.
        delete_other_zips(selected_paths)

        if UPLOAD_DB:
            upload_list.append(DB_PATH)

        if ENABLE_FTP and upload_list:
            try:
                upload_to_ftp(upload_list)
            except Exception as e:
                log(f"[ERROR] FTP upload failed: {e}")
                traceback.print_exc()
                log("The database was still created successfully at: " + DB_PATH)
        else:
            log("FTP disabled/skipped. Database ready at: " + DB_PATH)
    finally:
        conn.close()


if __name__ == "__main__":
    main()
