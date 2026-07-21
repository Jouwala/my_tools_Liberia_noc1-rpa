"""
数据清洗脚本（保留格式版）：
1. 读取 outband_details / inventory_details / inband_details 三个表
2. 删除 Warehouse Name == "Liberia Monrovia material warehouse" 的行
3. 用 material_config.xlsx 填充 Material Name 空值（通过 Material No. 匹配）
4. 按 Material Type 列拆分为多个文件，保留原表格式
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

import os
from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.dirname(os.path.abspath(__file__))

SOURCE_FILES = [
    "outband_details.xlsx",
    "inventory_details.xlsx",
    "inband_details.xlsx",
]

WAREHOUSE_TO_REMOVE = "Liberia Monrovia material warehouse"


def build_material_name_mapping():
    """从 material_config.xlsx 构建 Material No. -> Material Name 映射"""
    config_path = os.path.join(DATA_DIR, "material_config.xlsx")
    if not os.path.exists(config_path):
        print("[WARN] material_config.xlsx 不存在，跳过 Material Name 填充")
        return {}

    wb = load_workbook(config_path, data_only=True)
    ws = wb.active

    # 找列索引
    header = [cell.value for cell in ws[1]]
    try:
        mat_no_idx = header.index("Material No.") + 1
        mat_name_idx = header.index("Material Name") + 1
    except ValueError as e:
        print(f"[WARN] material_config.xlsx 缺少必要列: {e}")
        return {}

    mapping = {}
    for row in ws.iter_rows(min_row=2, max_col=max(mat_no_idx, mat_name_idx)):
        mat_no = row[mat_no_idx - 1].value
        mat_name = row[mat_name_idx - 1].value
        if mat_no is not None and mat_name is not None:
            # 转为字符串统一处理
            mapping[str(mat_no).strip()] = mat_name

    wb.close()
    print(f"  material_config 映射表: {len(mapping)} 条记录")
    return mapping


def copy_cell_style(src_cell, dst_cell):
    """复制单元格样式"""
    dst_cell.font = copy(src_cell.font)
    dst_cell.fill = copy(src_cell.fill)
    dst_cell.border = copy(src_cell.border)
    dst_cell.alignment = copy(src_cell.alignment)
    dst_cell.number_format = src_cell.number_format


def copy_column_widths(src_ws, dst_ws):
    """复制列宽"""
    for col_idx, col_dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_idx].width = col_dim.width


def copy_row_heights(src_ws, dst_ws, src_row, dst_row):
    """复制行高"""
    if src_row in src_ws.row_dimensions:
        dst_ws.row_dimensions[dst_row].height = src_ws.row_dimensions[src_row].height


def process_file(filename: str, material_mapping: dict):
    filepath = os.path.join(DATA_DIR, filename)
    if not os.path.exists(filepath):
        print(f"[SKIP] 文件不存在: {filepath}")
        return

    base_name = os.path.splitext(filename)[0]
    print(f"\n{'='*60}")
    print(f"处理: {filename}")

    wb = load_workbook(filepath)
    ws = wb.active

    # 获取表头
    header = [cell.value for cell in ws[1]]

    # 找列索引 (1-based)
    try:
        warehouse_col = header.index("Warehouse Name") + 1
    except ValueError:
        print("  [WARN] 未找到 'Warehouse Name' 列，跳过")
        wb.close()
        return

    try:
        material_type_col = header.index("Material Type") + 1
    except ValueError:
        print("  [WARN] 未找到 'Material Type' 列，跳过")
        wb.close()
        return

    # Material No. 和 Material Name 列（可选）
    mat_no_col = header.index("Material No.") + 1 if "Material No." in header else None
    mat_name_col = header.index("Material Name") + 1 if "Material Name" in header else None

    # 收集数据行（按 Material Type 分组），同时过滤仓库
    grouped_rows = {}
    total_rows = 0
    removed_rows = 0
    filled_count = 0

    for row_idx in range(2, ws.max_row + 1):
        total_rows += 1
        warehouse_val = ws.cell(row=row_idx, column=warehouse_col).value

        # 过滤指定仓库
        if warehouse_val == WAREHOUSE_TO_REMOVE:
            removed_rows += 1
            continue

        material_type = ws.cell(row=row_idx, column=material_type_col).value
        if material_type is None:
            material_type = "Unknown"

        # 填充 Material Name 空值
        if mat_name_col and mat_no_col:
            mat_name_cell = ws.cell(row=row_idx, column=mat_name_col)
            if mat_name_cell.value is None or str(mat_name_cell.value).strip() == "":
                mat_no = ws.cell(row=row_idx, column=mat_no_col).value
                if mat_no is not None:
                    mat_no_str = str(mat_no).strip()
                    if mat_no_str in material_mapping:
                        mat_name_cell.value = material_mapping[mat_no_str]
                        filled_count += 1

        if material_type not in grouped_rows:
            grouped_rows[material_type] = []
        grouped_rows[material_type].append(row_idx)

    print(f"  原始行数: {total_rows}")
    print(f"  删除 '{WAREHOUSE_TO_REMOVE}' 行数: {removed_rows}")
    print(f"  Material Name 空值填充: {filled_count} 个")
    print(f"  Material Type 类别: {list(grouped_rows.keys())}")

    # 为每个 Material Type 创建新工作簿
    for mtype, row_indices in grouped_rows.items():
        new_wb = Workbook()
        new_ws = new_wb.active
        new_ws.title = ws.title

        # 复制表头（第1行）
        for col_idx in range(1, ws.max_column + 1):
            src_cell = ws.cell(row=1, column=col_idx)
            dst_cell = new_ws.cell(row=1, column=col_idx)
            dst_cell.value = src_cell.value
            copy_cell_style(src_cell, dst_cell)

        # 复制列宽
        copy_column_widths(ws, new_ws)

        # 复制数据行
        for new_row_idx, src_row_idx in enumerate(row_indices, start=2):
            copy_row_heights(ws, new_ws, src_row_idx, new_row_idx)
            for col_idx in range(1, ws.max_column + 1):
                src_cell = ws.cell(row=src_row_idx, column=col_idx)
                dst_cell = new_ws.cell(row=new_row_idx, column=col_idx)
                dst_cell.value = src_cell.value
                copy_cell_style(src_cell, dst_cell)

        # 保存
        out_filename = f"{base_name}_{mtype}.xlsx"
        out_path = os.path.join(DATA_DIR, out_filename)
        new_wb.save(out_path)
        new_wb.close()
        print(f"  -> 已保存: {out_filename} ({len(row_indices)} 行)")

    wb.close()


if __name__ == "__main__":
    print("构建 Material Name 映射...")
    material_mapping = build_material_name_mapping()

    for f in SOURCE_FILES:
        process_file(f, material_mapping)

    print("\n处理完成！")
