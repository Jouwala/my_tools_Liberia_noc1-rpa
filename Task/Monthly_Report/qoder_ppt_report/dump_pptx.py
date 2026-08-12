from pptx import Presentation
from pptx.util import Emu
p = Presentation(r'd:\01_RPA\Task\Monthly_Report\qoder_ppt_report\Orange Liberia.pptx')
print('slides:', len(p.slides), 'size:', p.slide_width, p.slide_height)
for i, s in enumerate(p.slides):
    print(f'===== SLIDE {i+1} layout={s.slide_layout.name} =====')
    for sh in s.shapes:
        t = sh.shape_type
        name = sh.name
        pos = f'L={sh.left} T={sh.top} W={sh.width} H={sh.height}' if sh.left is not None else ''
        txt = ''
        if sh.has_text_frame:
            txt = ' | '.join(r.text for para in sh.text_frame.paragraphs for r in para.runs)
        print(f'  [{t}] {name} {pos}')
        if txt.strip():
            print(f'      TEXT: {txt[:300]}')
        if sh.has_table:
            tb = sh.table
            print(f'      TABLE {len(tb.rows)}x{len(tb.columns)}')
            for ri, row in enumerate(tb.rows):
                cells = [c.text.replace(chr(10),'/') for c in row.cells]
                print(f'        R{ri}: {cells}')
        if sh.has_chart:
            ch = sh.chart
            print(f'      CHART type={ch.chart_type}')
            try:
                cats = list(ch.plots[0].categories)
                print(f'        cats({len(cats)}): {cats[:5]}...{cats[-3:] if len(cats)>5 else ""}')
            except Exception as e:
                print('        cats err', e)
            for ser in ch.series:
                try:
                    vals = list(ser.values)
                    print(f'        series "{ser.name}": {vals[:5]}...{vals[-3:] if len(vals)>5 else ""}')
                except Exception as e:
                    print('        series err', e)
