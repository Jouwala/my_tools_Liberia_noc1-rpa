import openpyxl
wb = openpyxl.load_workbook(r'D:\01_RPA\Task\my_app\data\ms_sla\ZTE OLR MS Contract KPIs & WOs Achievement Auto-Reports.xlsx', data_only=True)
print('sheets:', wb.sheetnames)
ws = wb['Contract KPIs Report (New)']
print('dims:', ws.dimensions, 'max_row:', ws.max_row, 'max_col:', ws.max_column)
for row in ws.iter_rows(min_row=1, max_row=8, max_col=min(ws.max_column, 40)):
    vals = [f'{c.coordinate}={c.value}' for c in row if c.value is not None]
    if vals:
        print(vals)
