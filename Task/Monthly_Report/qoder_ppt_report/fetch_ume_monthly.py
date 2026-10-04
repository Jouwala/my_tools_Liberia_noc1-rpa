# -*- coding: utf-8 -*-
"""Orange Liberia 月报数据拉取（独立任务）

从 RAN UME 网关一次性拉取过去 12 个月【月粒度】NW 级 KPI 数据（GSM/UMTS/FDD），
供 Orange Liberia.pptx 第二页（Network KPIs 达标统计）与第三页（Traffic 趋势）使用。

- 不复用、不写入 my_app 的 lbr_ms_sla_kpi.sqlite3，数据落在本工作区 data/ 目录
- 登录流程（OAuth2 + RSA）与查询/导出协议参考 my_app backend/services/ms_sla_sync.py
- 输出：
    data/exports/*.xlsx                     原始导出文件（留档）
    data/orange_liberia_monthly.sqlite3     解析入库（GSM/UMTS/FDD_NW_1month）

用法：
    python fetch_ume_monthly.py [--months 12]
"""
import argparse
import base64
import gzip
import io
import json
import logging
import os
import random
import re
import socket
import sqlite3
import ssl
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


# --- 凭据引导（2026-10-04）：口令不写在源码里，统一走共享入口 ---
# 优先环境变量，其次「设置 → 凭据管理」页面（credential_store）。
import os as _os, sys as _sys
# 注意：判据是 credentials.py 这个**文件**，不是 shared/ 这个**目录**。
# Task/Unified_Tools/shared 也存在但里面没有 credentials.py，只判目录会
# 命中它、import 失败后被 except 吞掉 -> 静默降级为空口令（实测踩过）。
_ws = _os.path.dirname(_os.path.abspath(__file__))
_cand = None
while _ws:
    _try = _os.path.join(_ws, "Unified_Tools", "shared", "credentials.py")
    if _os.path.isfile(_try):
        _cand = _try
        break
    _p = _os.path.dirname(_ws)
    if _p == _ws:
        break
    _ws = _p
if _cand:
    _sys.path.insert(0, _os.path.dirname(_os.path.dirname(_cand)))
try:
    from shared.credentials import get_credential as _get_credential
except Exception:  # 凭据库不可用时降级，不让脚本直接崩
    def _get_credential(_key, env=None, default=""):
        return (_os.environ.get(env, "") if env else "") or default

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("fetch_ume_monthly")

# ============================================================
# 配置
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
EXPORT_DIR = DATA_DIR / "exports"
DB_PATH = DATA_DIR / "orange_liberia_monthly.sqlite3"

UME_HOST = "10.100.28.27"
UME_PORT = 28001
UME_USERNAME = "NOC"
UME_PASSWORD = _get_credential("zte_ume", env="RAN_UME_PASS")

HTTP_MAX_RETRIES = 3
HTTP_RETRY_DELAY = 1.0
HTTP_CONNECT_TIMEOUT = 15.0
HTTP_READ_TIMEOUT = 120.0

QUERY_TIMEOUT = 900      # 单次查询轮询超时（秒）
QUERY_POLL_INTERVAL = 5  # 轮询间隔（秒）

# my_app auth_manager 的鉴权缓存（可选复用，TTL 内免重复登录）
AUTH_CACHE_FILE = Path(r"D:\01_RPA\Task\auth_cache\10.100.28.27_28001.json")
AUTH_CACHE_TTL = 1700

SLA_HISTORY_PATH = "/api/pm-analysis/v2/history"
SLA_HISTORY_POLL_PATH = "/api/pm-analysis/v1/history"
SLA_HISTORY_EXPORT_PATH = "/api/pm-analysis/v3/history/export"
SLA_REFERER = f"https://{UME_HOST}:{UME_PORT}/iui/pm/angularTask/dist/pmManagement.html"

GR_MONTH = "2592000"  # 1Month 粒度
GROUP_LAYER_NW = "all"

# 查询配置（与 my_app SLA_QUERY_CONFIGS 保持一致）
SLA_QUERY_CONFIGS = {
    "GSM": {
        "templateid": "1638004472426",
        "templatename": "KPI_RPA_SITE_MS_SLA_GSM",
        "motid": "gsmPm.cell",
        "nfctid": "MRNC",
        "containvcell": 0,
        "indexfillingtype": "zeroDenominator",
        "showbusinesscolumn": False,
        "versiondatacompare": False,
        "items": [
            "306024", "900241213", "901415213", "306002", "306006", "306003",
            "306008", "306035", "306036", "900008213", "1675910745777", "306004",
            "1638004469689", "1675910745799",
        ],
        "itemsposition": [0] * 14,
        "itemsposition1": [0] * 14,
    },
    "UMTS": {
        "templateid": "1638004472427",
        "templatename": "KPI_RPA_SITE_MS_SLA_UMTS",
        "motid": "wPm.UtranCell,wPm.LogicCell",
        "nfctid": "MRNC",
        "containvcell": 2,
        "indexfillingtype": "zeroDenominator",
        "showbusinesscolumn": False,
        "versiondatacompare": False,
        "items": [
            "300249", "300240", "300294", "300243", "901593214", "300020",
            "300046", "300073", "300476", "300478", "300532", "300857",
            "300824", "300840", "901597214", "901598214", "300139", "900005214",
            "901411214", "901410214", "901620214", "901621214", "901622214",
            "901623214", "901604214", "901605214", "901603214", "901596214",
        ],
        "itemsposition": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0, 1, 1, 0, 0, 0, 0, 0, 0, 0, 0],
        "itemsposition1": [0] * 28,
    },
    "FDD": {
        "templateid": "1638004472428",
        "templatename": "KPI_RPA_SITE_MS_SLA_FDD",
        "motid": "EUtranCellFDD",
        "businessmotid": "EUtranCellFDD,FDDV3.EUtranCellFDD",
        "nfctid": "ITBBU,SDR,#",
        "containvcell": 0,
        "indexfillingtype": "zeroDenominator",
        "showbusinesscolumn": True,
        "items": [
            "1638004469333", "1638004469335", "1638004469337", "1638004469340",
            "1638004469341", "1638004469330", "1638004469331", "C373343602",
            "C373343605", "1638004469342", "1638004469343", "1638004469344",
            "1638004469345", "C373200029", "C373364118", "C373353961",
            "1638004469332", "1638004469785", "1603456136636", "1638004477799",
        ],
        "itemsposition": [0] * 20,
        "itemsposition1": [2, 2, 2, 2, 2, 2, 2, 1, 1, 2, 2, 2, 2, 1, 1, 1, 2, 2, 2, 2],
        "itempositionbusiness": [0, 0, 0, 0, 0, 0, 0, 1, 1, 0, 0, 0, 0, 1, 1, 1, 0, 0, 0, 0],
    },
}

# 第二页 Network KPIs 必备列（对应 Contract KPIs Report (New) K4:Q）
REQUIRED_KPI_COLUMNS = {
    "GSM": ["OCI_2G_CSSR_CS(%)", "TCH in call drop rate(include handover)(%)"],
    "UMTS": [
        "U31_3G Call Setup Success Rate - CS_901593_2_UROP.RNC-MO",
        "Ratio of CS Call Drop(%)",
        "U31_3G+ Call Setup Success Rate - HSDPA Data services_901597_2_UROP.RNC-MO",
        "U31_3G+ Drop Call Rate - HSDPA Data services_901604_2_UROP.RNC-MO",
        "Cell Inter-RAT CS Outgoing Handover Success Rate(WCDMA->GSM)(%)",
    ],
    "FDD": ["[FDD]E-RAB Drop Rate[SDR_ITBBU]"],
}
# 第三页 Traffic 必备列
REQUIRED_TRAFFIC_COLUMNS = {
    "GSM": ["TCH total traffic number(erl)",
            "U31_ZTE_Traffic of Uplink Data TBF (GByte)_900241_2_gv4.bsc-MO",
            "U31_ZTE_Traffic of Downlink Data TBF (GByte)_901415_2_gv4.bsc-MO"],
    "UMTS": ["Cell Traffic Volume, CS(Erl)",
             "U31_Cell DL Traffic Volume, PS(GB)_901411_2_UROP.RNC-MO",
             "U31_Cell UL Traffic Volume, PS(GB)_901410_2_UROP.RNC-MO"],
    "FDD": ["[FDD]DL Traffic Volume (GBytes)-new", "[FDD]UL Traffic Volume (GBytes)-new"],
}


# ============================================================
# HTTP 基础（单线程简化版，参考 ms_sla_sync._http_request）
# ============================================================
_conn = {"sock": None}


def _create_ssl_socket():
    ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    ctx.set_ciphers("ALL:@SECLEVEL=0")
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    sock.settimeout(HTTP_CONNECT_TIMEOUT)
    ssock = ctx.wrap_socket(sock)
    ssock.connect((UME_HOST, UME_PORT))
    ssock.settimeout(HTTP_READ_TIMEOUT)
    return ssock


def _get_socket():
    if _conn["sock"] is None:
        _conn["sock"] = _create_ssl_socket()
    return _conn["sock"]


def _close_socket():
    ssock = _conn["sock"]
    _conn["sock"] = None
    if ssock is not None:
        try:
            ssock.close()
        except OSError:
            pass


def _decode_chunked(data):
    result = bytearray()
    pos = 0
    while pos < len(data):
        line_end = data.find(b"\r\n", pos)
        if line_end == -1:
            break
        try:
            chunk_size = int(data[pos:line_end].decode("utf-8", errors="replace").strip(), 16)
        except ValueError:
            break
        if chunk_size == 0:
            break
        pos = line_end + 2
        result.extend(data[pos:pos + chunk_size])
        pos += chunk_size + 2
    return bytes(result)


def _http_request(method, path, headers_extra=None, body=None, retry=0):
    """原始 HTTPS 请求，返回 (status_line, resp_headers, raw_header_text, body_bytes)。"""
    try:
        ssock = _get_socket()
    except (ssl.SSLError, socket.timeout, ConnectionError, OSError) as e:
        if retry < HTTP_MAX_RETRIES:
            delay = HTTP_RETRY_DELAY * (2 ** retry) + random.uniform(0, 0.25)
            logger.warning(f"[HTTP] 连接失败({retry + 1}): {e}，{delay:.1f}s 后重试")
            time.sleep(delay)
            return _http_request(method, path, headers_extra=headers_extra, body=body, retry=retry + 1)
        raise

    headers = (
        f"{method} {path} HTTP/1.1\r\n"
        f"Host: {UME_HOST}:{UME_PORT}\r\n"
        f"Accept: application/json, text/plain, */*\r\n"
        f"Accept-Language: en-US,en;q=0.9\r\n"
        f"Accept-Encoding: identity\r\n"
        f"Connection: keep-alive\r\n"
        f"User-Agent: Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36\r\n"
    )
    if headers_extra:
        headers += headers_extra
    if body:
        headers += f"Content-Length: {len(body)}\r\n"
    headers += "\r\n"

    request_bytes = headers.encode("utf-8")
    if body:
        request_bytes += body if isinstance(body, bytes) else body.encode("utf-8")

    try:
        ssock.sendall(request_bytes)
    except (socket.timeout, ConnectionError, OSError) as e:
        _close_socket()
        if retry < HTTP_MAX_RETRIES:
            delay = HTTP_RETRY_DELAY * (2 ** retry) + random.uniform(0, 0.25)
            logger.warning(f"[HTTP] 发送失败({retry + 1}): {e}，{delay:.1f}s 后重试")
            time.sleep(delay)
            return _http_request(method, path, headers_extra=headers_extra, body=body, retry=retry + 1)
        raise

    response_data = bytearray()
    content_length = None
    header_end_found = False
    header_end_pos = 0

    while True:
        try:
            chunk = ssock.recv(65536)
            if not chunk:
                break
            response_data.extend(chunk)
        except socket.timeout:
            break

        if not header_end_found:
            pos = response_data.find(b"\r\n\r\n")
            if pos != -1:
                header_end_found = True
                header_end_pos = pos
                hp = response_data[:pos].decode("utf-8", errors="replace")
                for line in hp.split("\r\n"):
                    if line.lower().startswith("content-length:"):
                        try:
                            content_length = int(line.split(":", 1)[1].strip())
                        except ValueError:
                            pass
                        break

        if header_end_found and content_length is not None:
            if len(response_data) - (header_end_pos + 4) >= content_length:
                break

    if header_end_found and content_length is not None:
        body_data = response_data[header_end_pos + 4:]
        if len(body_data) < content_length:
            _close_socket()
            raise ConnectionError(f"响应截断: {len(body_data)}/{content_length} bytes")

    header_end = response_data.find(b"\r\n\r\n")
    if header_end == -1:
        _close_socket()
        if retry < HTTP_MAX_RETRIES:
            delay = HTTP_RETRY_DELAY * (2 ** retry) + random.uniform(0, 0.25)
            logger.warning(f"[HTTP] 响应无法解析({retry + 1})，{delay:.1f}s 后重试")
            time.sleep(delay)
            return _http_request(method, path, headers_extra=headers_extra, body=body, retry=retry + 1)
        raise Exception("无法解析 HTTP 响应")

    header_part = response_data[:header_end].decode("utf-8", errors="replace")
    resp_body = bytes(response_data[header_end + 4:])
    status_line = header_part.split("\r\n")[0]

    resp_headers = {}
    for line in header_part.split("\r\n")[1:]:
        if ": " in line:
            k, v = line.split(": ", 1)
            resp_headers[k.lower()] = v

    if "chunked" in resp_headers.get("transfer-encoding", ""):
        resp_body = _decode_chunked(resp_body)
        _close_socket()
    elif "close" in resp_headers.get("connection", "").lower() or content_length is None:
        _close_socket()

    if "gzip" in resp_headers.get("content-encoding", ""):
        resp_body = gzip.decompress(resp_body)

    return status_line, resp_headers, header_part, resp_body


def _rsa_encrypt_pem(plaintext, public_key_b64):
    from cryptography.hazmat.primitives.asymmetric import padding as asym_padding
    from cryptography.hazmat.primitives import serialization
    pem = f"-----BEGIN PUBLIC KEY-----\n{public_key_b64}\n-----END PUBLIC KEY-----"
    public_key = serialization.load_pem_public_key(pem.encode())
    encrypted = public_key.encrypt(plaintext.encode("utf-8"), asym_padding.PKCS1v15())
    return base64.b64encode(encrypted).decode("ascii")


# ============================================================
# 登录（参考 auth_manager._do_login，另可复用 my_app 鉴权缓存）
# ============================================================
def _try_cached_auth():
    """尝试复用 my_app auth_manager 的鉴权缓存（TTL 内有效）。"""
    try:
        if not AUTH_CACHE_FILE.exists():
            return None
        stat = AUTH_CACHE_FILE.stat()
        if time.time() - stat.st_mtime > AUTH_CACHE_TTL:
            return None
        data = json.loads(AUTH_CACHE_FILE.read_text(encoding="utf-8"))
        cookie = data.get("cookie")
        token = data.get("forgery_token")
        if cookie and token:
            logger.info("[AUTH] 复用 my_app 鉴权缓存")
            return {"cookie": cookie, "forgery_token": token}
    except Exception as e:
        logger.warning(f"[AUTH] 缓存读取失败: {e}")
    return None


def ume_login(max_retries=3):
    """ZTE 网管 OAuth2 RSA 登录，返回 {"cookie","forgery_token"}。"""
    cached = _try_cached_auth()
    if cached:
        return cached

    last_error = None
    for attempt in range(1, max_retries + 1):
        try:
            if attempt > 1:
                time.sleep(2 ** (attempt - 1))

            # Step 1: 授权页，获取初始 Cookie 与 loginParameters
            status, hdrs, raw_hdrs, body = _http_request(
                "GET", "/api/oauth2/v1/authorize?scope=user.login&response_type=code"
                       "&redirect_uri=/uportal/framework/default.html")
            init_cookies = []
            forgery_token = ""
            for line in raw_hdrs.split("\r\n"):
                if line.lower().startswith("set-cookie:"):
                    cookie_val = line.split(": ", 1)[1].split(";")[0]
                    if cookie_val.split("=")[1]:
                        init_cookies.append(cookie_val)
                        if "forgerydefense" in cookie_val:
                            forgery_token = cookie_val.split("=", 1)[1]
            init_cookie_str = "; ".join(init_cookies)

            html = body.decode("utf-8", errors="replace")
            lp_match = re.search(r'var\s+loginParameters\s*=\s*["\']([^"\']+)["\']', html)
            login_parameters = lp_match.group(1) if lp_match else ""

            # Step 2: RSA 公钥
            rsa_headers = f"Cookie: {init_cookie_str}\r\nX-Requested-With: XMLHttpRequest\r\n"
            status, hdrs, raw_hdrs, body = _http_request(
                "GET", "/api/oauth2/v1/login/publickey", headers_extra=rsa_headers)
            public_key_b64 = ""
            if "200" in status:
                public_key_b64 = json.loads(body.decode("utf-8")).get("rsaPublicKey", "")
            if not public_key_b64:
                pk_match = re.search(r'var\s+publicKey\s*=\s*["\']([^"\']+)["\']', html)
                if pk_match:
                    public_key_b64 = pk_match.group(1)
            if not public_key_b64:
                raise Exception("无法获取 RSA 公钥")

            # Step 3/4: 加密密码并提交登录
            encrypted_pwd = _rsa_encrypt_pem(UME_PASSWORD, public_key_b64)
            login_payload = json.dumps({
                "username": UME_USERNAME,
                "password": encrypted_pwd,
                "parameters": login_parameters,
            })
            login_headers = (
                f"Content-Type: application/json\r\n"
                f"Cookie: {init_cookie_str}\r\n"
                f"Referer: https://{UME_HOST}:{UME_PORT}/api/oauth2/v1/authorize?"
                f"scope=user.login&response_type=code&redirect_uri=/uportal/framework/default.html\r\n"
                f"X-Requested-With: XMLHttpRequest\r\n"
            )
            if forgery_token:
                login_headers += f"forgerydefense: {forgery_token}\r\n"

            status, hdrs, raw_hdrs, body = _http_request(
                "POST", "/api/oauth2/v1/login_rsa",
                headers_extra=login_headers, body=login_payload)

            cookies = list(init_cookies)
            for line in raw_hdrs.split("\r\n"):
                if line.lower().startswith("set-cookie:"):
                    cookie_val = line.split(": ", 1)[1].split(";")[0]
                    if cookie_val.split("=")[1]:
                        cookie_name = cookie_val.split("=")[0]
                        cookies = [c for c in cookies if not c.startswith(cookie_name + "=")]
                        cookies.append(cookie_val)
                        if "forgerydefense" in cookie_val:
                            forgery_token = cookie_val.split("=", 1)[1]

            if not any("Z-AUTH-CODE" in c for c in cookies):
                raise Exception("登录失败 - 未获取到 Z-AUTH-CODE cookie")

            logger.info(f"[AUTH] 登录成功 ({len(cookies)} cookies)")
            return {"cookie": "; ".join(cookies), "forgery_token": forgery_token}

        except Exception as e:
            last_error = e
            _close_socket()
            logger.warning(f"[AUTH] 登录尝试 {attempt} 失败: {e}")

    raise Exception(f"UME 登录最终失败: {last_error}")


def _sla_auth_headers(cookie, forgery_token, extra=None):
    headers = (
        f"Content-Type: application/json; charset=UTF-8\r\n"
        f"Accept: application/json, text/plain, */*\r\n"
        f"Accept-Encoding: gzip\r\n"
        f"Cookie: {cookie}\r\n"
        f"Referer: {SLA_REFERER}\r\n"
        f"X-Requested-With: XMLHttpRequest\r\n"
        f"language-option: en-US\r\n"
        f"timeZoneInfo: Default,Atlantic/Reykjavik,\r\n"
    )
    if forgery_token:
        headers += f"forgerydefense: {forgery_token}\r\n"
    if extra:
        for k, v in extra.items():
            headers += f"{k}: {v}\r\n"
    return headers


# ============================================================
# 查询 / 轮询 / 导出（参考 ms_sla_sync）
# ============================================================
def _build_payload(tech, start_iso, end_iso):
    cfg = SLA_QUERY_CONFIGS[tech]
    items = cfg["items"]
    item_count = len(items)
    timeinfo = {
        "mode": 1, "start": start_iso, "end": end_iso, "relative": -1,
        "timezonetype": "1", "clienttimezone": "Atlantic/Reykjavik",
        "showtimezoneoffset": False, "querylongtime": {"type": 3, "value": 1},
        "timespanfilter": "00:00:00-24:00:00", "dayfiltertype": 0,
        "dayinmonthfilter": "1111111111111111111111111111111",
    }
    payload = {
        "arraycounterinfos": [], "datatype": 0, "displaymode": 0,
        "favoriteconditionname": cfg["templatename"], "gr": GR_MONTH,
        "groupattributetype": 0, "groupbygroup": False, "groupbyplmn": False,
        "grouptype": 0, "items": items,
        "itemsposition": cfg.get("itemsposition", [0] * item_count),
        "itemsposition1": cfg.get("itemsposition1", [0] * item_count),
        "motid": cfg["motid"], "nfctid": cfg["nfctid"],
        "plmnpdfdisplaymode": False, "queryallbusiness": False,
        "showbusinesscolumn": cfg["showbusinesscolumn"], "showintegrity": False,
        "straightquery": False, "templatename": cfg["templatename"],
        "templateid": cfg["templateid"],
    }
    if "versiondatacompare" in cfg:
        payload["versiondatacompare"] = cfg["versiondatacompare"]
    payload.update({"needdeployfilter": False, "timeinfo": timeinfo,
                    "indexfillingtype": cfg["indexfillingtype"]})
    if "businessmotid" in cfg:
        payload["businessmotid"] = cfg["businessmotid"]
    if "itempositionbusiness" in cfg:
        payload["itempositionbusiness"] = cfg["itempositionbusiness"]
    payload.update({
        "thresholdinfos": [], "templatetype": 0, "filterlayer": "all",
        "grouplayer": GROUP_LAYER_NW, "containvcell": cfg["containvcell"],
        "resfilterattribute": "", "uilocations": [], "reslocation": [], "inexistenceres": False,
    })
    if "businessmotid" in cfg:
        payload["groupbusinessids"] = []
        payload["businessfilterscope"] = []
    payload.update({"firstDayOfWeek": 7, "iswireless": True, "hidehistorystopbutton": True,
                    "opencustomizehistoryexport": False, "opentemplatecustomcolumn": False})
    return payload


def _create_query(cookie, forgery_token, payload):
    body = json.dumps(payload, ensure_ascii=False)
    headers = _sla_auth_headers(cookie, forgery_token)
    status, hdrs, raw_hdrs, resp_body = _http_request(
        "POST", SLA_HISTORY_PATH, headers_extra=headers, body=body.encode("utf-8"))
    if "200" not in status:
        raise Exception(f"创建查询失败: {status} {resp_body[:200]!r}")
    data = json.loads(resp_body.decode("utf-8", errors="replace"))
    queryid = data.get("queryid") or data.get("id")
    if not queryid:
        raise Exception(f"未返回 queryid: {data}")
    return queryid


def _wait_query_ready(cookie, forgery_token, queryid):
    path = f"{SLA_HISTORY_POLL_PATH}?type=0&queryid={queryid}"
    headers = _sla_auth_headers(cookie, forgery_token,
                                extra={"Accept": "application/json, text/plain, */*"})
    deadline = time.time() + QUERY_TIMEOUT
    while time.time() < deadline:
        status, hdrs, raw_hdrs, resp_body = _http_request("GET", path, headers_extra=headers)
        data = json.loads(resp_body.decode("utf-8", errors="replace"))
        progress = data.get("progress", 0)
        querystate = data.get("querystate", 0)
        if querystate == 1 or progress >= 100:
            return data
        time.sleep(QUERY_POLL_INTERVAL)
    raise Exception(f"查询超时（{QUERY_TIMEOUT}s）")


def _export_query(cookie, forgery_token, queryid, filename) -> bytes:
    payload = {
        "id": queryid, "filename": filename, "chartlayout": 1, "contentmode": 0,
        "customizefieldtype": "0", "exportfiletype": 1, "exportserialnumber": False,
        "exporttableheaders": False, "formatwithts": 0, "imageinfo": [],
        "itemshowtype": 0, "tableheaders": {},
    }
    body = json.dumps(payload, ensure_ascii=False)
    headers = _sla_auth_headers(cookie, forgery_token)
    status, hdrs, raw_hdrs, resp_body = _http_request(
        "POST", SLA_HISTORY_EXPORT_PATH, headers_extra=headers, body=body.encode("utf-8"))
    if "200" not in status:
        raise Exception(f"导出失败: {status} {resp_body[:200]!r}")
    return resp_body


# ============================================================
# 解析与入库
# ============================================================
def read_xlsx_bytes(file_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter)
        headers = [str(col).strip() if col is not None else f"col_{i}"
                   for i, col in enumerate(header_row)]
        data = [list(row) for row in rows_iter if not all(v is None for v in row)]
    finally:
        wb.close()
    return headers, data


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def write_to_sqlite(db_path: Path, table_name: str, headers: list, data: list):
    """整表重建写入（独立任务，无需增量 upsert）。"""
    con = sqlite3.connect(str(db_path))
    try:
        cur = con.cursor()
        cur.execute(f"DROP TABLE IF EXISTS {_quote_ident(table_name)}")
        cols_sql = ", ".join(_quote_ident(h) for h in headers)
        cur.execute(f"CREATE TABLE {_quote_ident(table_name)} ({cols_sql})")
        placeholders = ", ".join("?" * len(headers))
        cur.executemany(
            f"INSERT INTO {_quote_ident(table_name)} VALUES ({placeholders})", data)
        con.commit()
        return len(data)
    finally:
        con.close()


def fetch_monthly(months: int):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    # 时间窗口：覆盖过去 months 个整月（起点取月初）
    now = datetime.now(timezone.utc)
    y, m = now.year, now.month
    for _ in range(months):
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    start_iso = datetime(y, m, 1, tzinfo=timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    end_iso = now.strftime("%Y-%m-%dT%H:%M:%S.000Z")
    logger.info(f"查询窗口: {start_iso} ~ {end_iso}（月粒度，目标 {months} 个月）")

    auth = ume_login()
    cookie, forgery_token = auth["cookie"], auth["forgery_token"]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary = {}
    for tech in SLA_QUERY_CONFIGS:
        templatename = SLA_QUERY_CONFIGS[tech]["templatename"]
        logger.info(f"===== [{tech}] 开始查询（1Month / NW）=====")
        payload = _build_payload(tech, start_iso, end_iso)
        queryid = _create_query(cookie, forgery_token, payload)
        logger.info(f"[{tech}] Query ID: {queryid}")
        status = _wait_query_ready(cookie, forgery_token, queryid)
        logger.info(f"[{tech}] 查询完成: totalcount={status.get('totalcount')}")

        filename = f"{templatename}_NW_1Month"
        xlsx_bytes = _export_query(cookie, forgery_token, queryid, filename)
        xlsx_path = EXPORT_DIR / f"{filename}_{stamp}.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)
        logger.info(f"[{tech}] 已导出: {xlsx_path.name} ({len(xlsx_bytes)} bytes)")

        headers, data = read_xlsx_bytes(xlsx_bytes)
        table_name = f"{tech}_NW_1month"
        rows = write_to_sqlite(DB_PATH, table_name, headers, data)
        months_got = sorted({str(r[0])[:7] for r in data}) if data else []
        logger.info(f"[{tech}] 入库 {table_name}: {rows} 行，月份覆盖: {months_got}")

        missing = [c for c in REQUIRED_KPI_COLUMNS[tech] + REQUIRED_TRAFFIC_COLUMNS[tech]
                   if c not in headers]
        if missing:
            logger.warning(f"[{tech}] 缺少必备列: {missing}")
        summary[tech] = {"rows": rows, "months": months_got, "missing_columns": missing,
                         "xlsx": str(xlsx_path)}

    _close_socket()
    return summary


def main():
    parser = argparse.ArgumentParser(description="拉取 RAN UME 月粒度 NW KPI 数据")
    parser.add_argument("--months", type=int, default=12, help="回溯月数（默认 12）")
    args = parser.parse_args()

    t0 = time.time()
    summary = fetch_monthly(args.months)

    logger.info("=" * 60)
    logger.info(f"完成，耗时 {time.time() - t0:.1f}s")
    ok = True
    for tech, info in summary.items():
        flag = "OK" if info["rows"] > 0 and not info["missing_columns"] else "WARN"
        if info["rows"] == 0:
            ok = False
        logger.info(f"  [{flag}] {tech}: {info['rows']} 行，{len(info['months'])} 个月"
                    f"（{info['months'][0] if info['months'] else '-'} ~ "
                    f"{info['months'][-1] if info['months'] else '-'}），"
                    f"缺失列 {len(info['missing_columns'])}")
    logger.info(f"数据库: {DB_PATH}")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
