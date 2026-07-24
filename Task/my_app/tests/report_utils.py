"""Shared test-record writer: produces Markdown + JSON (+ Excel) test records.

A "record" is a plain dict:
{
  "project": str,
  "suite": "backend_api" | "frontend_e2e",
  "started_at": ISO str,
  "ended_at": ISO str,
  "environment": {...},
  "summary": {"total": int, "passed": int, "failed": int, "skipped": int},
  "cases": [ case, ... ]
}
A "case" dict:
{
  "id": str, "module": str, "name": str,
  "method": str, "path": str,
  "expected": str, "actual_status": int|None,
  "latency_ms": float|None,
  "passed": bool, "skipped": bool,
  "assertions": [str, ...],      # human-readable assertion lines
  "error": str|None,
  "extra": {...},                # arbitrary detail (response summary, screenshot path...)
}
"""
from __future__ import annotations
import json
import os
from datetime import datetime, timezone


def _now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def make_record(project: str, suite: str, environment: dict) -> dict:
    return {
        "project": project,
        "suite": suite,
        "started_at": _now_iso(),
        "ended_at": None,
        "environment": environment,
        "summary": {"total": 0, "passed": 0, "failed": 0, "skipped": 0},
        "cases": [],
    }


def finalize_record(record: dict) -> dict:
    total = passed = failed = skipped = 0
    for c in record["cases"]:
        if c.get("skipped"):
            skipped += 1
        elif c.get("passed"):
            passed += 1
        else:
            failed += 1
        total += 1
    record["summary"] = {
        "total": total,
        "passed": passed,
        "failed": failed,
        "skipped": skipped,
        "pass_rate": round(passed / total * 100, 1) if total else 0.0,
    }
    record["ended_at"] = _now_iso()
    return record


def _status_badge(c: dict) -> str:
    if c.get("skipped"):
        return "SKIP"
    return "PASS" if c.get("passed") else "FAIL"


def write_markdown(record: dict, path: str) -> None:
    s = record["summary"]
    lines = []
    lines.append(f"# 测试记录 · {record['project']}")
    lines.append("")
    lines.append(f"- **测试套件**: `{record['suite']}`")
    lines.append(f"- **开始时间**: {record.get('started_at')}")
    lines.append(f"- **结束时间**: {record.get('ended_at')}")
    env = record.get("environment", {})
    if env:
        lines.append("- **测试环境**:")
        for k, v in env.items():
            lines.append(f"  - {k}: `{v}`")
    lines.append("")
    lines.append("## 汇总")
    lines.append("")
    lines.append(f"| 总计 | 通过 | 失败 | 跳过 | 通过率 |")
    lines.append(f"|------|------|------|------|--------|")
    lines.append(f"| {s['total']} | {s['passed']} | {s['failed']} | {s['skipped']} | {s.get('pass_rate', 0)}% |")
    lines.append("")

    # group by module
    modules: dict[str, list] = {}
    for c in record["cases"]:
        modules.setdefault(c.get("module", "其他"), []).append(c)

    for mod, cases in modules.items():
        lines.append(f"## 模块：{mod}")
        lines.append("")
        lines.append("| 用例ID | 名称 | 方法 | 路径 | 期望 | 实际状态 | 耗时(ms) | 结果 |")
        lines.append("|--------|------|------|------|------|----------|-----------|------|")
        for c in cases:
            lines.append(
                f"| {c.get('id')} | {c.get('name')} | {c.get('method') or '-'} | "
                f"`{c.get('path') or '-'}` | {c.get('expected') or '-'} | "
                f"{c.get('actual_status') if c.get('actual_status') is not None else '-'} | "
                f"{c.get('latency_ms') if c.get('latency_ms') is not None else '-'} | "
                f"**{_status_badge(c)}** |"
            )
        lines.append("")

    # detailed sections for failures / all cases with extra
    lines.append("## 详细记录")
    lines.append("")
    for c in record["cases"]:
        badge = _status_badge(c)
        lines.append(f"### {c.get('id')} · {c.get('name')} — **{badge}**")
        lines.append("")
        meta = []
        if c.get("method"):
            meta.append(f"- 方法/路径: `{c.get('method')} {c.get('path')}`")
        if c.get("expected"):
            meta.append(f"- 期望: {c.get('expected')}")
        if c.get("actual_status") is not None:
            meta.append(f"- 实际状态: {c.get('actual_status')}")
        if c.get("latency_ms") is not None:
            meta.append(f"- 耗时: {c.get('latency_ms')} ms")
        if c.get("extra"):
            for k, v in c["extra"].items():
                meta.append(f"- {k}: {v}")
        if meta:
            lines.extend(meta)
            lines.append("")
        if c.get("assertions"):
            lines.append("- 断言:")
            for a in c["assertions"]:
                lines.append(f"  - {a}")
            lines.append("")
        if c.get("error"):
            lines.append(f"- **错误**: {c.get('error')}")
            lines.append("")
        if c.get("skipped"):
            lines.append("- 说明: 该用例跳过（如依赖外部网络/凭据，需手动验证）。")
            lines.append("")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def write_json(record: dict, path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(record, f, ensure_ascii=False, indent=2)


def write_excel(record: dict, path: str) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    s = record["summary"]
    ws.append(["测试套件", record["suite"]])
    ws.append(["项目", record["project"]])
    ws.append(["开始", record.get("started_at")])
    ws.append(["结束", record.get("ended_at")])
    ws.append([])
    ws.append(["总计", "通过", "失败", "跳过", "通过率"])
    ws.append([s["total"], s["passed"], s["failed"], s["skipped"], f"{s.get('pass_rate',0)}%"])
    for cell in ws[6]:
        cell.font = Font(bold=True)

    ws2 = wb.create_sheet("用例明细")
    headers = ["用例ID", "模块", "名称", "方法", "路径", "期望", "实际状态",
               "耗时(ms)", "结果", "断言", "错误", "截图/备注"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(vertical="center")
    for c in record["cases"]:
        ws2.append([
            c.get("id"), c.get("module"), c.get("name"), c.get("method") or "",
            c.get("path") or "", c.get("expected") or "",
            c.get("actual_status") if c.get("actual_status") is not None else "",
            c.get("latency_ms") if c.get("latency_ms") is not None else "",
            _status_badge(c),
            "\n".join(c.get("assertions", [])),
            c.get("error") or "",
            (c.get("extra") or {}).get("screenshot") or (c.get("extra") or {}).get("response_summary") or "",
        ])
    # widths
    widths = [10, 14, 22, 8, 34, 18, 10, 10, 8, 50, 30, 40]
    for i, w in enumerate(widths, start=1):
        ws2.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    ws2.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return True


def write_reports(record: dict, out_dir: str, stem: str) -> dict:
    md = os.path.join(out_dir, f"{stem}.md")
    js = os.path.join(out_dir, f"{stem}.json")
    xlsx = os.path.join(out_dir, f"{stem}.xlsx")
    write_markdown(record, md)
    write_json(record, js)
    excel_ok = write_excel(record, xlsx)
    return {"markdown": md, "json": js, "excel": xlsx if excel_ok else None}
