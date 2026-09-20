# -*- coding: utf-8 -*-
"""MS SLA Contract KPIs 报表「日期行滚动」回归测试。

## 背景（2026-09-19 线上故障）
`ms_sla_full`（MS SLA 完整同步）连续三次失败于：

    AttributeError: property 'customHeight' of 'RowDimension' object has no setter

现场分支：报表文件里已有 9/1~9/17 共 17 行日期行，当天目标是 9/1~9/18 共 18 行
→ `_roll_date_rows` 走 `insert_rows` 分支 → `_copy_row_style` 里
`dst_dim.customHeight = src_dim.customHeight` 赋值只读 property → 整个任务失败。

openpyxl 里 `RowDimension.customHeight` 是只读 property（值 = `ht is not None`），
只应设置 `height`，`customHeight` 会自动跟随。本测试锁死该行为，防止回归。

运行：
    cd D:/01_RPA/Task/my_app
    python -m pytest tests/test_ms_sla_report_row_roll.py -q
"""
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

# 让测试能 import backend 包
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from backend.services.ms_sla_report import (  # noqa: E402
    TARGET_FILE,
    TARGET_SHEET,
    _copy_row_style,
    _locate_report_rows,
    _roll_date_rows,
    _target_report_dates,
)

MONTH_ROW = 6
FIRST_DATE_ROW = 7
TEMPLATE_ROWS = 17
TEMPLATE_LAST_HEIGHT = 21.5
FILL_RGB = "FFEEEEEE"


def _build_ws(n_rows):
    """构造与线上同构的最小工作表：1 个 Month 行 + n_rows 个连续日期行（带样式）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TARGET_SHEET

    ws.cell(row=MONTH_ROW, column=1, value="Month").font = Font(bold=True)
    for c in range(2, 29):
        ws.cell(row=MONTH_ROW, column=c, value=f"KPI {c}")

    base = datetime(2026, 9, 1)
    for i in range(n_rows):
        r = FIRST_DATE_ROW + i
        cell = ws.cell(row=r, column=1, value=base + timedelta(days=i))
        cell.font = Font(bold=(i % 2 == 0))
        cell.fill = PatternFill("solid", start_color=FILL_RGB)
        for c in range(2, 29):
            ws.cell(row=r, column=c, value=i + c / 100.0)
    # 样板行（最后一行）设置行高，验证 insert 后行高被复制
    ws.row_dimensions[FIRST_DATE_ROW + n_rows - 1].height = TEMPLATE_LAST_HEIGHT
    return wb, ws


def _sep_dates(n):
    return [datetime(2026, 9, 1) + timedelta(days=i) for i in range(n)]


def test_roll_inserts_rows_and_copies_style():
    """现场分支：17 行 → 18 行（insert_rows + _copy_row_style 必须不抛异常）。"""
    wb, ws = _build_ws(TEMPLATE_ROWS)
    try:
        month_row, date_rows = _locate_report_rows(ws)
        assert month_row == MONTH_ROW
        assert len(date_rows) == TEMPLATE_ROWS

        new_rows = _roll_date_rows(ws, month_row, date_rows, _sep_dates(18))

        assert len(new_rows) == 18
        assert [ws.cell(row=r, column=1).value.strftime("%Y-%m-%d") for r in new_rows] == \
            [f"2026-09-{d:02d}" for d in range(1, 19)]
        last = new_rows[-1]
        # 新增行的格式（填充/加粗/行高）必须跟着复制，否则条件格式与边框掉
        assert ws.cell(row=last, column=1).fill.start_color.rgb == FILL_RGB
        assert ws.cell(row=last, column=1).font.bold is True
        assert ws.row_dimensions[last].height == TEMPLATE_LAST_HEIGHT
        # customHeight 是派生值，设了 height 后自动为 True
        assert ws.row_dimensions[last].customHeight is True
    finally:
        wb.close()


def test_roll_deletes_rows():
    """31 行 → 18 行：删除分支同样要写对日期且不抛异常。"""
    wb, ws = _build_ws(31)
    try:
        month_row, date_rows = _locate_report_rows(ws)
        new_rows = _roll_date_rows(ws, month_row, date_rows, _sep_dates(18))

        assert len(new_rows) == 18
        assert ws.cell(row=new_rows[-1], column=1).value.strftime("%Y-%m-%d") == "2026-09-18"
        assert ws.max_row == MONTH_ROW + 18
    finally:
        wb.close()


def test_roll_keeps_rows_when_count_unchanged():
    wb, ws = _build_ws(TEMPLATE_ROWS)
    try:
        month_row, date_rows = _locate_report_rows(ws)
        new_rows = _roll_date_rows(ws, month_row, date_rows, _sep_dates(TEMPLATE_ROWS))

        assert new_rows == date_rows
        assert ws.cell(row=new_rows[-1], column=1).value.strftime("%Y-%m-%d") == "2026-09-17"
    finally:
        wb.close()


def test_custom_height_is_read_only_in_this_openpyxl():
    """把故障根因写进测试：该 property 无 setter，谁再赋值谁失败。"""
    wb, ws = _build_ws(2)
    try:
        dim = ws.row_dimensions[FIRST_DATE_ROW]
        with pytest.raises(AttributeError):
            dim.customHeight = True
    finally:
        wb.close()


def test_copy_row_style_only_sets_height():
    """_copy_row_style 复制行高但不触碰只读 property。"""
    wb, ws = _build_ws(3)
    try:
        src, dst = FIRST_DATE_ROW, FIRST_DATE_ROW + 10
        ws.row_dimensions[src].height = TEMPLATE_LAST_HEIGHT

        _copy_row_style(ws, src, dst)  # 修复前此处抛 AttributeError

        assert ws.row_dimensions[dst].height == TEMPLATE_LAST_HEIGHT
        assert ws.row_dimensions[dst].customHeight is True
        assert ws.cell(row=dst, column=1).fill.start_color.rgb == FILL_RGB
    finally:
        wb.close()


def test_real_report_file_can_be_rolled_to_18_days():
    """用线上真实报表文件（只读加载、不保存）跑一遍现场日期区间。"""
    if not Path(TARGET_FILE).exists():
        pytest.skip(f"线上报表文件不存在: {TARGET_FILE}")

    wb = openpyxl.load_workbook(TARGET_FILE)
    try:
        ws = wb[TARGET_SHEET]
        month_row, date_rows = _locate_report_rows(ws)
        if month_row is None or not date_rows:
            pytest.skip("报表文件缺少 Month 行或日期行，跳过")
        if date_rows != list(range(date_rows[0], date_rows[0] + len(date_rows))):
            pytest.skip("报表文件日期行不连续，_roll_date_rows 走只改写分支，跳过")

        dates = _target_report_dates(datetime(2026, 9, 19))  # 9/1 ~ 9/18
        new_rows = _roll_date_rows(ws, month_row, date_rows, dates)

        assert len(new_rows) == len(dates) == 18
        assert ws.cell(row=new_rows[-1], column=1).value.strftime("%Y-%m-%d") == "2026-09-18"
        assert isinstance(ws.cell(row=new_rows[0], column=1).value, (datetime, date))
    finally:
        wb.close()
