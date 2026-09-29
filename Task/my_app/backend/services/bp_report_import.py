# -*- coding: utf-8 -*-
"""BP Report —— Excel 一次性导入。

把 `data/umlaut_data/BP Report.xlsx`（10 sheet / 14.3MB）导入 `data/bp_report.sqlite3`。

=========================== 为什么需要「导入」而不是每次读 Excel ===========================

1. `BP Cells List` 有 **42.4 万行**，直接读一次要几十秒且吃内存；导入后所有查询都是
   SQL 索引扫描。
2. Excel 的 6 个汇总 sheet 是**上游（umlaut）算好的结果**，本项目拿不到其原始明细；
   把它们落库才能与 RAN_UME 重算值**逐日对账**（见 `bp_report_sync.py`）。
3. 导入是**一次性**的，重复执行幂等（全部 UPSERT），因此可以随时重跑。

=========================== 脏值清洗（基于 42.4 万行实测） ===========================

`BP Cells List` 有两个列的取值被污染，实测分布：

- `Cluster` 列混入约 **100 个纯数字值**（`'4593'` / `'9103'` / `'618-07'` …），
  实为 Site ID / eNodeB ID **错填**（列错位）。典型错行（sheet 第 168855 行起）：
  `(2025-03-03, '5593', 'Yekepa3-S1', '4593', 'Rural')` —— 倒数第二列是 ME ID、最后一列才是真 Cluster。
- `Layer` 列混入 `Rural` / `ERA` / `Cluster*` 等**非层值**（合计约 250 行）。

处理原则（`bp_caliber` 里实现，此处只调用）：
- Cluster：合法值原样保留，其余一律归 `Unknown` 并计数 —— **不做数值猜测**。
- Layer：列值合法则以列值为准；非法时按小区名后缀回推（`-F`→L800 / `-S`→L1800 20M /
  `-L`→L1800 15M / `-X`→L2600）；仍失败归 `Unknown`。

**为什么 Cluster 不回推到 `Layer` 列**：错行里 `Layer` 刚好放着真 Cluster，看似可以
两列对调修复。但实测只有 **7 行左右**符合「Cluster 是数字 且 Layer 是合法 Cluster」这一
形态，且无法排除其他行是别种错位；为避免引入**未被证实的**映射，统一归 `Unknown`，
把计数暴露在 `/meta`，由人工决定是否补录。

=========================== 字典的两种来源 ===========================

`bp_cells`  sheet 只给出「**是 BP 的那些天**」的小区，因此：

- `bp_cell_dict` = 出现在 `BP Cells List` 的全部小区（4269 个），**只含曾经是 BP 的小区**。
- 非 BP 小区的 Cluster/Layer 无从得知（Excel 没有全量小区清单）。
  → 重算时，字典外的小区归 `Unknown` 并在 `/meta` 暴露（`bp_report_sync.py` 处理）。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from backend.services import bp_caliber as cal
from backend.services import bp_report_store as store

logger = logging.getLogger("bp_report_import")

SERVICE_DIR = Path(__file__).resolve().parent
DEFAULT_XLSX = SERVICE_DIR.parent.parent / "data" / "umlaut_data" / "BP Report.xlsx"

#: 逐批写入的行数。42.4 万行不能一次攒在内存里（AGENTS.md 原则 6 的边界处理）。
BATCH_SIZE = 20000


def _to_int(value: Any) -> Optional[int]:
    """宽松整数转换；空/非法返回 None（不返回 0 —— 0 与「没数据」必须区分）。"""
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _to_float(value: Any) -> Optional[float]:
    """宽松浮点转换；空/非法返回 None。"""
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _rows(ws) -> Iterable[tuple]:
    """跳过表头逐行产出。用 read_only 迭代器，不把整表读进内存。"""
    it = ws.iter_rows(values_only=True)
    next(it, None)
    for row in it:
        if row and any(v is not None and v != "" for v in row):
            yield row


# ============================================================
# 各 sheet 的解析
# ============================================================

def _import_bp_cells(ws, conn, now: str) -> Dict[str, Any]:
    """`BP Cells List` → `bp_cell_daily` + `bp_cell_dict`。

    这是唯一给出「小区 × 日期 × 是否 BP」明细的 sheet，也是 Cluster/Layer 字典的来源。
    列序：Date | ManagedElement ID | E-UTRAN FDD Cell Name | Cluster | Layer
    """
    detail: List[tuple] = []
    # 字典聚合：cell_name → {me_id, cluster 投票, layer 投票, first/last day, 天数}
    agg: Dict[str, Dict[str, Any]] = {}
    dirty_cluster = 0
    dirty_layer = 0
    total = 0

    for row in _rows(ws):
        day = cal.day_str(row[0])
        cell = (str(row[2]).replace("\xa0", " ").strip() if row[2] is not None else "")
        if not day or not cell:
            continue
        me_id = str(row[1]).strip() if row[1] is not None else None
        raw_cluster, raw_layer = row[3], row[4]

        if cal.is_dirty_cluster(raw_cluster):
            dirty_cluster += 1
        if cal.is_dirty_layer(raw_layer):
            dirty_layer += 1

        cluster = cal.normalize_cluster(raw_cluster)
        layer = cal.normalize_layer(raw_layer, cell)
        total += 1

        detail.append((day, cell, me_id, cluster, layer,
                       1, None, None, None, None, None,
                       cluster, "excel", now))

        a = agg.setdefault(cell, {"me_id": me_id, "first": day, "last": day, "days": 0,
                                  "clusters": {}, "layers": {}})
        if me_id and not a["me_id"]:
            a["me_id"] = me_id
        a["first"] = min(a["first"], day)
        a["last"] = max(a["last"], day)
        a["days"] += 1
        a["clusters"][cluster] = a["clusters"].get(cluster, 0) + 1
        a["layers"][layer] = a["layers"].get(layer, 0) + 1

        if len(detail) >= BATCH_SIZE:
            store.upsert_many(conn, "bp_cell_daily",
                              ("day", "cell_name", "me_id", "cluster", "layer", "is_bp",
                               "prb_dl", "thr_dl", "active_user", "dl_gb", "ul_gb",
                               "region", "source", "updated_at"),
                              detail, ("day", "cell_name"))
            detail.clear()

    if detail:
        store.upsert_many(conn, "bp_cell_daily",
                          ("day", "cell_name", "me_id", "cluster", "layer", "is_bp",
                           "prb_dl", "thr_dl", "active_user", "dl_gb", "ul_gb",
                           "region", "source", "updated_at"),
                          detail, ("day", "cell_name"))

    # 字典：一个小区多天出现时，取**出现次数最多**的 Cluster/Layer（众数），
    # 因为它反映的是该小区的稳定归属；出现次数相等时取字典序较小者，保证可重复。
    dict_rows = []
    for cell, a in agg.items():
        cluster = max(sorted(a["clusters"]), key=lambda k: a["clusters"][k], default=cal.UNKNOWN)
        layer = max(sorted(a["layers"]), key=lambda k: a["layers"][k], default=cal.UNKNOWN)
        dict_rows.append((cell, a["me_id"], cluster, layer,
                          a["first"], a["last"], a["days"], now))
    store.upsert_many(conn, "bp_cell_dict",
                      ("cell_name", "me_id", "cluster", "layer",
                       "first_day", "last_day", "bp_days", "updated_at"),
                      dict_rows, ("cell_name",))

    return {"rows": total, "cells": len(agg), "dirty_cluster": dirty_cluster,
            "dirty_layer": dirty_layer}


def _import_overall_daily(ws, conn, now: str) -> int:
    """`Overall - Daily` → `bp_overall_daily`。

    列序：Begin Time | Date(Mon..Sun) | Number of BP cells | Total Cells | BP ratio |
         4G_Data_Traffic(TB) - excluding ERA
    """
    rows = []
    for row in _rows(ws):
        day = cal.day_str(row[0])
        if not day:
            continue
        bp_cells = _to_int(row[2])
        rows.append((day, bp_cells, _to_int(row[3]), _to_float(row[4]), _to_float(row[5]),
                     bp_cells,  # bp_cells_excel：Excel 原生值，重算时不覆盖
                     str(row[1]).strip() if row[1] is not None else None,
                     "excel", now))
    store.upsert_many(conn, "bp_overall_daily",
                      ("day", "bp_cells", "total_cells", "bp_ratio", "traffic_tb",
                       "bp_cells_excel", "native_date", "source", "updated_at"),
                      rows, ("day",))
    return len(rows)


def _import_regional_daily(ws, conn, now: str) -> int:
    """`Regional - Daily` → `bp_region_daily`（同时也是 bp_cell_daily.region 的对账源）。

    列序：Begin Time | Region | BP Number | Total Cell | BP Ratio | Total Traffic (TB)
    """
    rows = []
    for row in _rows(ws):
        day = cal.day_str(row[0])
        region = str(row[1]).strip() if row[1] is not None else ""
        if not day or not region:
            continue
        rows.append((day, region, _to_int(row[2]), _to_int(row[3]),
                     _to_float(row[4]), _to_float(row[5]), "excel", now))
    store.upsert_many(conn, "bp_region_daily",
                      ("day", "region", "bp_number", "total_cell", "bp_ratio",
                       "traffic_tb", "source", "updated_at"),
                      rows, ("day", "region"))
    return len(rows)


def _import_layer_daily(ws, conn, now: str) -> int:
    """`Frequency layer - Daily` → `bp_layer_daily`。

    列序：Begin Time | Region | 每层三列 (BP Number / Total Cell / BP Ratio)。
    层的顺序**写死在表头**里，不假设固定的 4 层 —— 逐列解析表头拿层名，
    这样将来 Excel 增删层（例如补上 `L1800 10M`）时不用改代码。
    表头形如 `L1800 20M BP Number` / `L1800 20M Total Cell` / `L1800 20M BP Ratio`。
    """
    it = ws.iter_rows(values_only=True)
    header = next(it, None) or ()
    # 解析表头：列号 → (layer, 字段)
    col_map: Dict[int, Tuple[str, str]] = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        text = str(h).replace("\xa0", " ").strip()
        for suffix, field in ((" BP Number", "bp"), (" Total Cell", "total"),
                              (" BP Ratio", "ratio")):
            if text.endswith(suffix):
                layer = text[: -len(suffix)].strip()
                col_map[i] = (layer, field)
                break

    # (day, region, layer) → 行字典，逐行合并三列
    buffer: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    written = 0
    for row in it:
        if not row or len(row) < 2:
            continue
        day = cal.day_str(row[0])
        region = str(row[1]).strip() if row[1] is not None else ""
        if not day or not region:
            continue
        for i, (layer, field) in col_map.items():
            if i >= len(row):
                continue
            v = row[i]
            if v is None or v == "":
                continue
            key = (day, region, layer)
            slot = buffer.setdefault(key, {"bp": None, "total": None, "ratio": None})
            slot[field] = _to_int(v) if field != "ratio" else _to_float(v)

        if len(buffer) >= BATCH_SIZE:
            written += len(buffer)
            _flush_layer(conn, buffer, now)
    written += len(buffer)
    _flush_layer(conn, buffer, now)
    return written


def _flush_layer(conn, buffer: Dict[Tuple[str, str, str], Dict[str, Any]], now: str) -> None:
    if not buffer:
        return
    rows = [(d, r, l, s["bp"], s["total"], s["ratio"], "excel", now)
            for (d, r, l), s in buffer.items()]
    store.upsert_many(conn, "bp_layer_daily",
                      ("day", "region", "layer", "bp_number", "total_cell",
                       "bp_ratio", "source", "updated_at"),
                      rows, ("day", "region", "layer"))
    buffer.clear()


def _import_overall_week(ws, conn, now: str) -> int:
    """`Overall - 1 Weekly` → `bp_overall_week`。

    列序：Week(`W48, 2023`) | BP cells | Total cells | BP Ratio | Total Traffic (TB)

    `week` 存 ISO 周键。**Excel 末尾 39 行是裸周号 `W1`..`W39` 不带年份**（实测：
    带年份的到 `W52, 2025` 为止），因此年份要靠 `cal.resolve_excel_weeks()` 按
    「周号回绕即跨年」推断，不能逐行独立解析。
    周粒度没有逐日明细，`bp_days`/`total_days` 留空 —— 这两个中间量只由重算侧填。
    """
    records: List[Tuple[str, Any]] = []
    for row in _rows(ws):
        native = str(row[0]).strip() if row[0] is not None else ""
        if not native:
            continue
        records.append((native, row))

    resolved = cal.resolve_excel_weeks([n for n, _ in records])
    rows = []
    for (native, row), (y, w) in zip(records, resolved):
        week = f"{y}-W{w:02d}" if w else native
        bp_cells = _to_int(row[1])
        rows.append((week, None, None, None, bp_cells, _to_int(row[2]),
                     _to_float(row[3]), _to_float(row[4]), native, "excel", now))
    store.upsert_many(conn, "bp_overall_week",
                      ("week", "bp_days", "total_days", "is_bp_week", "bp_cells",
                       "total_cells", "bp_ratio", "traffic_tb", "native_period",
                       "source", "updated_at"),
                      rows, ("week",))
    return len(rows)


def _import_overall_month(ws, conn, now: str) -> int:
    """`Overall - Monthly` → `bp_overall_month`。

    列序：Month(`Apr, 2023`) | Area(`Whole Network`) | BP Ratio | Total Traffic [TB] |
          （两个无名列，实测是 BP cells 与 Total cells —— **列名缺失，靠位置取**）

    ⚠ 这张表**没有给出 BP cells / Total cells 的列名**（表头第 5、6 列为空），
    实测取值形如 `(493, 5938)`，与同月 BP ratio 一致（493/5938≈0.0830），
    因此按位置 4、5 读取。若上游改版导致错位，`/meta` 的月对账会立刻暴露。
    """
    rows = []
    for row in _rows(ws):
        native = str(row[0]).strip() if row[0] is not None else ""
        if not native:
            continue
        month = cal.parse_excel_month(native) or native
        # 位置 4/5 是 BP cells / Total cells；ratio 在位置 2
        bp_cells = _to_int(row[4]) if len(row) > 4 else None
        total_cells = _to_int(row[5]) if len(row) > 5 else None
        rows.append((month, None, None, None, bp_cells, total_cells,
                     _to_float(row[2]), _to_float(row[3]), native, "excel", now))
    store.upsert_many(conn, "bp_overall_month",
                      ("month", "bp_weeks", "total_weeks", "is_bp_month", "bp_cells",
                       "total_cells", "bp_ratio", "traffic_tb", "native_period",
                       "source", "updated_at"),
                      rows, ("month",))
    return len(rows)


#: sheet 名 → 解析函数。只列**作为独立数据源**的 6 张主表；
#: `Sheet1`/`Sheet3` 是透视残留（见 bp_caliber.SHEET_IGNORED），不导入。
SHEET_HANDLERS = (
    (cal.SHEET_BP_CELLS, _import_bp_cells),
    (cal.SHEET_OVERALL_DAILY, _import_overall_daily),
    (cal.SHEET_REGIONAL_DAILY, _import_regional_daily),
    (cal.SHEET_LAYER_DAILY, _import_layer_daily),
    (cal.SHEET_OVERALL_WEEK, _import_overall_week),
    (cal.SHEET_OVERALL_MONTH, _import_overall_month),
)


def run_import(path: Optional[Path] = None, rebuild: bool = True) -> Dict[str, Any]:
    """执行一次性导入。

    Args:
        path:    Excel 路径，默认 `data/umlaut_data/BP Report.xlsx`。
        rebuild: 默认**先删库重建**。看似比 UPSERT 粗暴，但这是必要的：主键
                 （如 `bp_overall_week.week`）本身由解析逻辑产生，一旦解析规则变化
                 （例如把裸周号 `W9` 修成 `2026-W09`），UPSERT 只会新增行而**留下
                 旧键的孤儿行**，页面就会出现重复周期。导入是纯全量动作，重算代价
                 （约 40 秒）远低于脏数据的排查代价。

    Returns:
        dict: {success, sheets, rows, elapsed_seconds, errors}
    """
    import openpyxl

    t0 = time.time()
    xlsx = Path(path) if path else DEFAULT_XLSX
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    result: Dict[str, Any] = {"success": False, "sheets": {}, "rows": 0,
                              "elapsed_seconds": 0, "errors": []}

    if not xlsx.exists():
        result["errors"].append(f"Excel 不存在: {xlsx}")
        result["elapsed_seconds"] = round(time.time() - t0, 2)
        return result

    if rebuild and store.db_exists():
        # 只删本模块自己的物化库，绝不碰 db/app.db 等其它库
        try:
            os.remove(store.DB_PATH)
            logger.info("[BP] 已删除旧库 %s，全量重建", store.DB_PATH)
        except OSError as e:
            result["errors"].append(f"删除旧库失败: {e}")
            result["elapsed_seconds"] = round(time.time() - t0, 2)
            return result

    # data_only=True：取公式的**缓存值**。这些 sheet 大量使用透视/公式，
    # 取公式文本毫无意义（openpyxl 也拿不到计算结果）。
    # read_only=True：42.4 万行的 sheet 必须流式读，否则内存直接爆。
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        with store.write_conn() as conn:
            store.ensure_schema(conn)
            for sheet_name, handler in SHEET_HANDLERS:
                if sheet_name not in wb.sheetnames:
                    result["errors"].append(f"sheet 缺失: {sheet_name}")
                    continue
                try:
                    out = handler(wb[sheet_name], conn, now)
                except Exception as e:  # 单表失败不放弃其余表
                    logger.exception("[BP] sheet %s 导入失败", sheet_name)
                    result["errors"].append(f"{sheet_name} 导入失败: {e}")
                    continue
                if isinstance(out, dict):
                    result["sheets"][sheet_name] = out
                    result["rows"] += out.get("rows", 0)
                else:
                    result["sheets"][sheet_name] = out
                    result["rows"] += out

            # 记录导入元信息
            store.set_meta(conn, "last_import_at", now, now)
            store.set_meta(conn, "source_file", str(xlsx), now)
            store.set_meta(conn, "source_mtime",
                           datetime.fromtimestamp(xlsx.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                           now)
            store.set_meta(conn, "caliber", cal.caliber_snapshot(), now)
    finally:
        wb.close()

    result["success"] = result["rows"] > 0 and not result["errors"]
    result["elapsed_seconds"] = round(time.time() - t0, 2)
    logger.info("[BP] 导入完成 rows=%s %.1fs errors=%s",
                result["rows"], result["elapsed_seconds"], result["errors"])
    return result


if __name__ == "__main__":
    import argparse

    logging.basicConfig(level=logging.INFO,
                        format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    ap = argparse.ArgumentParser(description="BP Report Excel 一次性导入")
    ap.add_argument("--file", default=None, help="Excel 路径（默认 data/umlaut_data/BP Report.xlsx）")
    args = ap.parse_args()
    print(json.dumps(run_import(Path(args.file) if args.file else None),
                     ensure_ascii=False, indent=2, default=str))
