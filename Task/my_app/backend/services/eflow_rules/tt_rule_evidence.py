# -*- coding: utf-8 -*-
"""TT 单规则命中证据（TT Rule Match Evidence）
=================================================
给「TT 快速分析 → 判定结论」提供**确定性**的规则命中证据，让结论不只是
纯文本推断，而是能落到「命中了哪条规则、规则来自哪个文件、规则谁维护、工单谁在处理」。

数据来源与口径（全部复用既有实现，**不产生第二套判定**）
------------------------------------------------------
* 规则匹配：`backend.services.eflow_rules.tt_batch_severity.check_batch()`
  —— 与每天 04:20 的定时任务 `eflow_tt_rules_check`、以及 Email Reports →
  `eFlow TT Rules Check` 页签里的矫正表**同源**（内部复用 `eflow_rules_matcher`），
  因此本模块给出的「命中规则」与那两处逐字一致。
* 规则明细：`eflow_TT_rules.xlsx`（TT_Rules 工作表）的
  规则类型 / 工单类型 / 触发步骤 / 优先级 / 严重级别 / 状态 / 创建人 / 更新人 / 时间 / 条件。
* 产物文件：`eflow_tt_rules_check_latest.json`（最近一次「收集 + 匹配」的时间与产物文件名）。

对外接口
--------
* `build_rule_evidence(tt_no, handler=None)` → dict（结构见函数 docstring）
* `render_evidence_text(evidence, lang)`   → str（喂给大模型的「本地规则引擎预判定」文本块）
* `render_local_conclusion(evidence, ...)`  → str（**程序化** Markdown 结论，无模型参与）

页面侧：`GET /api/eflow-qa/tt-ai-fast` 把本模块的结果挂在 `ai.rules` 上返回、
结论正文取 `ai.conclusion`，前端在「判定结论（本地规则引擎）」卡片里渲染成
「命中规则」明细块 + 结论 Markdown。`GET /api/eflow-qa/tt-ai`（模型路径）
仅在候选模型档全失败时才用同一份程序化结论兜底。
"""
from __future__ import annotations

import json
import logging
import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

# my_app 根目录：本文件 = <root>/backend/services/eflow_rules/tt_rule_evidence.py
_APP_ROOT = Path(__file__).resolve().parents[3]

# 与 eflow_tt_rules_check 定时任务、Email Reports 下载白名单共用的源 id
_SOURCE_ID = "eflow_tt_rules_check"
_DOWNLOAD_PREFIX = f"/api/email-reports/download/{_SOURCE_ID}"

# 规则表里要一并带出的列（页面与 AI 结论都要能说清「这条规则是什么、谁维护的」）
_RULE_COLUMNS = (
    "ID", "规则名称", "规则类型", "工单类型", "触发步骤", "描述",
    "优先级", "严重级别", "状态", "创建人", "更新人",
    "创建时间", "更新时间", "条件",
)

# 推单规则表（eflow_Push_Rule_Collector 采集落盘）里的列
_PUSH_RULE_COLUMNS = (
    "ID", "规则名称", "规则类型", "工单类型", "触发步骤", "描述",
    "优先级", "严重级别", "推送路径", "目标处理人", "状态",
    "创建人", "更新人", "创建时间", "更新时间", "动态项(JSON)", "条件",
)

# 规则表按 mtime 缓存：规则每天 04:20 刷新一次，Web 请求路径上没必要反复解析 xlsx
_RULE_ROWS_CACHE: Dict[str, Any] = {"mtime": None, "rows": {}}

# 推单规则表按 mtime 缓存（同上；与 _RULE_ROWS_CACHE 分开，两份文件刷新时间不同）
_PUSH_RULE_CACHE: Dict[str, Any] = {"mtime": None, "rows": []}

# 规则类型 → 展示名（eFlow 规则页签上的写法；未登记的按驼峰拆词兜底）
_RULE_TYPE_TEXT = {
    "ttPushOrderRule": "TT Push Order Rule",
    "ttLevelRule": "TT Level Rule",
    "ttDispatchRule": "TT Dispatch Rule",
    "ttGroupConfigRule": "TT Group Config Rule",
    "ttGroupDefaultRule": "TT Group Default Rule",
    "crGroupConfigRule": "CR Group Config Rule",
    "woGroupConfigRule": "WO Group Config Rule",
}

# 驼峰/下划线拆词时的首段大写规则（ttPushOrderRule → TT Push Order Rule）
_CAMEL_HEAD = {"tt": "TT", "wo": "WO", "cr": "CR", "bo": "BO", "ip": "IP", "fo": "FO"}

# eFlow 内部的 UUID 形式 id（分组 id 是 UUID，人名/账号才是可读 code）
_UUID_RE = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")

# 决定「FO 接单时是否自动建 WO / 推给谁」的触发步骤
# （与 tt_postmortem._analyze_dispatch_reason 的 target_step 保持同一口径）
_PUSH_DECISION_STEP = "FOAccept"



def _s(v: Any) -> str:
    """NaN / None / 'nan' → 空串；其余转字符串并去首尾空白（与 tt_batch_severity._s 同口径）。"""
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none", "nat", "null") else s


def _rules_xlsx() -> Path:
    """规则表路径（与 eflow_rules_matcher / openwebui_client 同一份配置）。"""
    from backend.services.eflow_rules import eflow_rules_matcher as M
    return M.RULES_XLSX


def _summary_json() -> Path:
    """最近一次「规则收集 + 匹配检查」的摘要（定时任务 eflow_tt_rules_check 落盘）。"""
    return _rules_xlsx().parent / "eflow_tt_rules_check_latest.json"


def _load_rule_rows() -> Dict[str, Dict[str, str]]:
    """读规则表全部行（不过滤状态），按 mtime 缓存；返回 {规则ID: {列名: 值}}。"""
    from openpyxl import load_workbook

    path = _rules_xlsx()
    try:
        mtime = path.stat().st_mtime
    except OSError as e:
        raise RuntimeError(f"规则文件不可用：{path}（{e}）") from e

    if _RULE_ROWS_CACHE["mtime"] == mtime and _RULE_ROWS_CACHE["rows"]:
        return _RULE_ROWS_CACHE["rows"]

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "TT_Rules" not in wb.sheetnames:
            raise RuntimeError(f"规则表缺少 TT_Rules 工作表：{path}")
        ws = wb["TT_Rules"]
        rows = ws.iter_rows(values_only=True)
        header = [_s(h) for h in next(rows)]
        out: Dict[str, Dict[str, str]] = {}
        for vals in rows:
            rec = {k: _s(v) for k, v in zip(header, vals) if k}
            rid = rec.get("ID", "")
            if rid:
                out[rid] = rec
    finally:
        wb.close()
    _RULE_ROWS_CACHE.update({"mtime": mtime, "rows": out})
    return out


def _file_info(path: Path) -> Dict[str, Any]:
    """产物文件 → 展示用信息（相对路径 / 大小 / 修改时间 / 下载地址）。"""
    name = path.name
    info: Dict[str, Any] = {
        "name": name,
        "download_url": f"{_DOWNLOAD_PREFIX}/{name}",
        "exists": path.is_file(),
        "size": None,
        "mtime": None,
        "mtime_text": "",
    }
    try:
        rel = path.relative_to(_APP_ROOT).as_posix()
    except ValueError:
        rel = path.as_posix()
    info["rel_path"] = rel
    if info["exists"]:
        st = path.stat()
        info["size"] = st.st_size
        info["mtime"] = int(st.st_mtime)
        info["mtime_text"] = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(st.st_mtime))
    return info


def _rules_file_block() -> Dict[str, Any]:
    """规则表本身 + 最近一次检查产物的文件信息（供页面展示与下载）。"""
    xlsx = _rules_xlsx()
    block: Dict[str, Any] = {
        "xlsx": _file_info(xlsx),
        "report": _file_info(xlsx.parent / "MATCHING_REPORT.md"),
        "correction": _file_info(xlsx.parent / "eflow_tt_severity_correction.xlsx"),
        "task_id": _SOURCE_ID,
        "collected_at": "",
        "generated_at": "",
        "rules_count": None,
        "rules_updated_at": "",
        "rules_updated_ts": None,
    }
    try:
        st = xlsx.stat()
        block["rules_updated_at"] = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(st.st_mtime))
        block["rules_updated_ts"] = int(st.st_mtime)
    except OSError:
        pass
    try:
        summary = json.loads(_summary_json().read_text(encoding="utf-8"))
        block["collected_at"] = _s(summary.get("generated_at"))
        block["rules_count"] = summary.get("rules_count")
        rc = summary.get("rules_collected") or {}
        if rc.get("success") is False:
            block["collect_error"] = _s(rc.get("error"))
    except Exception as e:  # noqa: BLE001
        # 摘要缺失不影响主流程（首次部署或任务未跑过）
        logger.debug("读取 eflow_tt_rules_check 摘要失败（忽略）: %s", e)
    return block


def _matched_block(rule_id: str, match_details: str, priority: str,
                   expected_severity: str, changed: bool) -> Dict[str, Any]:
    """命中规则 → 规则表原始列 + 匹配详情（规则表里查不到时只给匹配结果）。"""
    rec: Dict[str, str] = {}
    try:
        rec = _load_rule_rows().get(rule_id, {})
    except Exception as e:  # noqa: BLE001
        logger.warning("读取规则明细失败（只返回匹配结果）: %s", e)
    return {
        "rule_id": rule_id,
        "rule_name": rec.get("规则名称", ""),
        "priority": priority or rec.get("优先级", ""),
        "severity": expected_severity or rec.get("严重级别", ""),
        "severity_changed": bool(changed),
        "rule_type": rec.get("规则类型", ""),
        "ticket_type": rec.get("工单类型", ""),
        "trigger_step": rec.get("触发步骤", ""),
        "description": rec.get("描述", ""),
        "status": rec.get("状态", ""),
        "creator": rec.get("创建人", ""),
        "updater": rec.get("更新人", ""),
        "created_at": rec.get("创建时间", ""),
        "updated_at": rec.get("更新时间", ""),
        "condition": rec.get("条件", ""),
        "match_details": match_details,
    }


def _push_rules_xlsx() -> Path:
    """推单规则表路径（与 eflow_TT_rules.xlsx 同目录，由 eflow_Push_Rule_Collector 落盘）。"""
    return _rules_xlsx().parent / "eflow_push_rules.xlsx"


def _load_push_rule_rows() -> List[Dict[str, str]]:
    """读推单规则表 Push_Rules 工作表全部行，按 mtime 缓存。"""
    from openpyxl import load_workbook

    path = _push_rules_xlsx()
    try:
        mtime = path.stat().st_mtime
    except OSError as e:
        raise RuntimeError(f"推单规则文件不可用：{path}（{e}）") from e

    if _PUSH_RULE_CACHE["mtime"] == mtime and _PUSH_RULE_CACHE["rows"]:
        return _PUSH_RULE_CACHE["rows"]

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Push_Rules" not in wb.sheetnames:
            raise RuntimeError(f"推单规则表缺少 Push_Rules 工作表：{path}")
        ws = wb["Push_Rules"]
        rows = ws.iter_rows(values_only=True)
        header = [_s(h) for h in next(rows)]
        out: List[Dict[str, str]] = []
        for vals in rows:
            rec = {k: _s(v) for k, v in zip(header, vals) if k}
            if rec.get("ID"):
                out.append(rec)
    finally:
        wb.close()
    _PUSH_RULE_CACHE.update({"mtime": mtime, "rows": out})
    return out


def _camel_text(raw: str) -> str:
    """驼峰/下划线标识 → 展示名（FOAccept → FO Accept；ttPushOrderRule → TT Push Order Rule）。"""
    s = _s(raw)
    if not s:
        return ""
    if s in _RULE_TYPE_TEXT:
        return _RULE_TYPE_TEXT[s]
    s = re.sub(r"[_\-.]+", " ", s)
    parts = re.findall(r"[A-Z]+(?=[A-Z][a-z])|[A-Z]?[a-z]+|[A-Z]+|\d+", s) or [s]
    if parts and parts[0] in _CAMEL_HEAD:
        parts[0] = _CAMEL_HEAD[parts[0]]
    return " ".join(parts)


def _pretty_push_path(raw: str) -> str:
    """推送路径 `FOAccept,2>FOProcess,8>FOProcess` → `FOAccept > FOProcess > FOProcess`。

    每段后面的 `,N` 是 eFlow 的步骤序号（内部实现细节），展示时去掉。
    """
    segs = []
    for seg in re.split(r"[>→]", _s(raw)):
        seg = re.sub(r",\s*\d+\s*$", "", seg.strip()).strip()
        if seg:
            segs.append(seg)
    return " > ".join(segs)


def _push_rule_dynamic(rec: Dict[str, str]) -> Dict[str, Any]:
    """解析 `动态项(JSON)` 列（推单路径 / 自动处理人 / 目标处理人的结构化原始值）。"""
    try:
        d = json.loads(rec.get("动态项(JSON)") or "{}")
        return d if isinstance(d, dict) else {}
    except Exception:  # noqa: BLE001
        return {}


def _items(raw: Any) -> List[Dict[str, str]]:
    """动态项里的 `[{id,type,name},...]` JSON 字符串 → 列表（解析失败返回空）。"""
    if not raw:
        return []
    try:
        items = json.loads(raw) if isinstance(raw, str) else raw
    except Exception:  # noqa: BLE001
        return []
    if not isinstance(items, list):
        return []
    return [it for it in items if isinstance(it, dict)]


def _auto_operator_text(dy: Dict[str, Any]) -> str:
    """自动处理人 → `FE Robot (fe_robot)`（eFlow 规则页签的写法：名称 + 账号）。"""
    parts = []
    for it in _items(dy.get("autoOperator")):
        name, uid = _s(it.get("name")), _s(it.get("id"))
        if name and uid:
            parts.append(f"{name} ({uid})")
        else:
            parts.append(name or uid)
    return "; ".join(p for p in parts if p)


def _target_handler_text(dy: Dict[str, Any], fallback: str = "") -> str:
    """目标处理人 → `ZTE_NOC_FO [ZTE_NOC_FO]`（eFlow 规则页签的写法：名称 + [编码]）。

    分组在 eFlow 里的 id 是 UUID（如 26833f43-…），不是人能读的编码；
    可读编码就是分组名本身（`ZTE_NOC_FO` / `3rd_Power` / `SOC-CEM`），
    所以 UUID 形式的 id 用名称代替 —— 与规则页签上看到的一致。
    """
    parts = []
    for it in _items(dy.get("targetHandler")):
        name, uid = _s(it.get("name")), _s(it.get("id"))
        code = uid if (uid and not _UUID_RE.match(uid)) else name
        if name and code:
            parts.append(f"{name} [{code}]")
        else:
            parts.append(name or code)
    if parts:
        return "; ".join(p for p in parts if p)
    return _s(fallback)   # 动态项缺失时退回 `目标处理人` 列（形如 ZTE_NOC_FO(group)）


def _push_rule_block(rec: Dict[str, str], match_details: str) -> Dict[str, Any]:
    """一条命中的推单规则 → 页面/结论要展示的全部字段。"""
    dy = _push_rule_dynamic(rec)
    return {
        "rule_id": _s(rec.get("ID")),
        "rule_name": _s(rec.get("规则名称")),
        "rule_type": _s(rec.get("规则类型")),
        "rule_type_text": _camel_text(rec.get("规则类型")),
        "ticket_type": _s(rec.get("工单类型")),
        "trigger_step": _s(rec.get("触发步骤")),
        "trigger_step_text": _camel_text(rec.get("触发步骤")),
        "priority": _s(rec.get("优先级")),
        "push_path": _pretty_push_path(rec.get("推送路径")),
        "auto_operator": _auto_operator_text(dy),
        "target_handler": _target_handler_text(dy, rec.get("目标处理人")),
        "status": _s(rec.get("状态")),
        "creator": _s(rec.get("创建人")),
        "updater": _s(rec.get("更新人")),
        "created_at": _s(rec.get("创建时间")),
        "updated_at": _s(rec.get("更新时间")),
        "description": _s(rec.get("描述")),
        "condition": _s(rec.get("条件")),
        "match_details": match_details,
    }


def _parse_push_conditions(cond_str: str) -> List[Dict[str, Any]]:
    """推单规则 `条件` 列 → 结构化条件（复用级别规则的 `parse_conditions`）。

    **两份表的条件列格式不同**，必须归一化后再解析：
      * 级别规则（`eflow_TT_rules.xlsx`）：`AND | eflow_ticket_ticket_wo.alarmCode in '9970, 9980' (string)`
        —— 连接词与字段之间有一个 `|` 分隔符；
      * 推单规则（`eflow_push_rules.xlsx`，由 `eflow_Push_Rule_Collector._fmt_rows` 生成）：
        `AND eflow_ticket_ticket_wo.alarmCode in '9970, 9980' (string)` —— **没有 `|`**。
    直接把后者喂给 `parse_conditions` 会得到空列表（它会跳过不含 `|` 的片段），
    表现为「所有推单规则都不命中」—— 这里补上分隔符，其余解析逻辑一字不改。
    """
    from backend.services.eflow_rules import eflow_rules_matcher as M

    txt = _s(cond_str)
    if not txt:
        return []
    parts: List[str] = []
    for raw in txt.split(";"):
        s = raw.strip()
        if not s:
            continue
        if "|" not in s:
            s = re.sub(r"^(AND|OR)\b", r"\1 |", s, count=1, flags=re.IGNORECASE)
            if "|" not in s:
                # 没有连接词的片段按 AND 处理（首行的连接词本来就被忽略）
                s = "AND | " + s
        parts.append(s)
    return M.parse_conditions(";".join(parts))


def build_push_rule_evidence(tt_no: str) -> Dict[str, Any]:
    """TT 单 → 命中的**推单规则**（TT Push Order Rule）证据。

    为什么要单独做：级别规则（ttLevelRule）只回答「这单该是什么级别」，
    而「这单会不会自动推 WO、推给谁、谁自动处理」由推单规则决定 ——
    页面「规则来源与责任人」必须把两者都摆出来，否则「责任人」是缺的。

    口径
    ----
    * 规则来源：`eflow_push_rules.xlsx`（`eflow_Push_Rule_Collector` 采集落盘，
      与 Email Reports → eFlow TT Rules Check 页签里可下载的那份**同一个文件**）。
    * 匹配：复用 `eflow_rules_matcher` 的 `parse_conditions` / `build_value_frame` /
      `build_rule_mask`，与级别规则**同一套求值代码**，只是换成推单规则表。
    * 排序：优先级升序 → 条件数多者优先 → ID 升序（与 `prepare_rules` 一致），
      取第一条为「实际生效」的推单规则，其余一并列出。
    * 字段局限：本地 `eflow_tt` 只有 alarmCode / vipSite / siteSubType(assetType) /
      netWorkType / downSiteCount 五个可评估字段；引用其他字段（category、
      previousOperationStep、alarmDurationTime 等）的规则**无法在本地判定**，
      单独计入 `undecidable` 并在界面上说明，不混进「未命中」。

    任何异常都不抛出：返回 available=False + error，由页面自行降级。
    """
    t0 = time.perf_counter()
    out: Dict[str, Any] = {
        "available": False,
        "error": "",
        "rules_file": {},
        "matched": [],
        "matched_count": 0,
        "other_steps": [],
        "other_steps_count": 0,
        "evaluated": 0,
        "undecidable": 0,
        "undecidable_rules": [],
        "fields": [],
        "elapsed_sec": 0.0,
    }
    try:
        path = _push_rules_xlsx()
        out["rules_file"] = _file_info(path)
        out["rules_file"]["download_url"] = f"{_DOWNLOAD_PREFIX}/{path.name}"
        out["rules_file"]["rules_count"] = None
        try:
            out["rules_file"]["mtime_text"] = time.strftime(
                "%Y-%m-%d %H:%M:%S", time.localtime(path.stat().st_mtime))
        except OSError:
            pass

        rows = [r for r in _load_push_rule_rows()
                if _s(r.get("规则类型")) == "ttPushOrderRule"
                and _s(r.get("状态")).upper() == "ACTIVE"]
        out["rules_file"]["rules_count"] = len(rows)

        from backend.services.eflow_rules import tt_batch_severity as batch
        from backend.services.eflow_rules import eflow_rules_matcher as M

        src_rows, _hit = batch.fetch_tt_raw([tt_no])
        if not src_rows:
            out["error"] = "本地 eflow_tt 未找到该工单，推单规则无法判定"
            return out

        import pandas as pd
        vf = M.build_value_frame(pd.DataFrame(src_rows))
        out["fields"] = [c for c in vf.columns]

        matched: List[Dict[str, Any]] = []
        undecidable: List[str] = []
        for rec in rows:
            conds = _parse_push_conditions(rec.get("条件") or "")
            if not conds:
                continue
            unknown = [c["field"] for c in conds if c["field"] not in vf.columns]
            if unknown:
                # 条件里含本地取不到的字段 → 本地判不了，不能算「未命中」
                undecidable.append(f"{rec.get('规则名称') or rec.get('ID')}"
                                   f"（{'/'.join(sorted(set(unknown)))}）")
                continue
            try:
                mask = M.build_rule_mask(vf, conds)
                if bool(mask.iloc[0]):
                    detail = _s(M.format_details(vf, conds).iloc[0])
                    matched.append(_push_rule_block(rec, detail))
            except Exception as e:  # noqa: BLE001
                logger.debug("推单规则 %s 求值失败（跳过）：%s", rec.get("ID"), e)
                continue

        def _sort_key(r: Dict[str, Any]):
            try:
                pri = int(float(r.get("priority") or 999))
            except (TypeError, ValueError):
                pri = 999
            return (pri, -len(r.get("match_details") or ""), _s(r.get("rule_id")))

        matched.sort(key=_sort_key)
        # 只有 `triggerStep=FOAccept` 的推单规则决定「FO 接单时是否自动建 WO / 推给谁」
        # （与 tt_postmortem._analyze_dispatch_reason 的 target_step 一致）；
        # 其他步骤（如 Confirm）命中的规则单列，避免把兜底规则混进「责任人」。
        at_step = [m for m in matched if _s(m.get("trigger_step")) == _PUSH_DECISION_STEP]
        others = [m for m in matched if _s(m.get("trigger_step")) != _PUSH_DECISION_STEP]

        out["matched"] = at_step
        out["matched_count"] = len(at_step)
        out["other_steps"] = others
        out["other_steps_count"] = len(others)
        out["evaluated"] = len(rows) - len(undecidable)
        out["undecidable"] = len(undecidable)
        out["undecidable_rules"] = undecidable[:10]
        out["available"] = True
        return out
    except Exception as e:  # noqa: BLE001
        out["error"] = f"推单规则匹配失败：{type(e).__name__}: {e}"
        logger.warning("TT 推单规则证据：%s", out["error"])
        return out
    finally:
        out["elapsed_sec"] = round(time.perf_counter() - t0, 2)


def build_rule_evidence(tt_no: str, handler: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """TT 单 → 规则命中证据（本地规则引擎权威判定 + 规则表元数据 + 产物文件）。

    参数
    ----
    tt_no   : TT 单号（大写，调用方已做白名单校验）
    handler : 工单处理人信息（可选），形如
              {"last_processer": ..., "last_processer_group": ..., "last_process_time": ...}
              —— 由调用方从 `inspect_tt` 报告的「派单」段提取，避免本模块再查一次库。

    返回
    ----
    {
      "available": True,                  # 规则引擎是否给出了可用判定
      "error": "",                        # available=False 时的原因（人读）
      "tt_no": "TT202609230017",
      "found": True,                      # 本地 eflow_tt 是否有该工单
      "level_status": "ok|adjust|unknown",
      "current_severity": "P3",
      "expected_severity": "P3",
      "matched": {规则明细 + match_details} | None,
      "rules_file": {...},                # eflow_TT_rules.xlsx + 产物文件信息
      "handler": {...},                   # 工单处理人（可能为空）
      "elapsed_sec": 0.4
    }

    任何异常都不会抛出：返回 available=False + error，由页面与 AI 提示自行降级。
    """
    t0 = time.perf_counter()
    out: Dict[str, Any] = {
        "available": False,
        "error": "",
        "tt_no": _s(tt_no),
        "found": False,
        "level_status": "unknown",
        "current_severity": "",
        "expected_severity": "",
        "matched": None,
        "rules_file": {},
        "push_rules": {},
        "handler": {k: _s(v) for k, v in (handler or {}).items() if _s(v)},
        "elapsed_sec": 0.0,
    }

    # 规则表 / 产物文件信息先取：即使匹配失败，页面也能显示「规则来自哪个文件」
    try:
        out["rules_file"] = _rules_file_block()
    except Exception as e:  # noqa: BLE001
        out["error"] = f"规则文件信息读取失败：{type(e).__name__}: {e}"
        logger.warning("TT 规则证据：%s", out["error"])

    # 推单规则（TT Push Order Rule）：与级别规则相互独立，先算，
    # 这样「级别规则未命中 / 本地无该单」时页面照样能说清「谁推单、推给谁」。
    try:
        out["push_rules"] = build_push_rule_evidence(out["tt_no"])
    except Exception as e:  # noqa: BLE001
        logger.warning("TT 推单规则证据失败（忽略）：%s", e)
        out["push_rules"] = {"available": False, "error": f"{type(e).__name__}: {e}",
                             "matched": [], "matched_count": 0}

    try:
        from backend.services.eflow_rules import tt_batch_severity as batch

        res = batch.check_batch([out["tt_no"]], limit=1)
        rows = res.get("rows") or []
        row = rows[0] if rows else None
        if not row or not row.get("found"):
            out["error"] = out["error"] or "本地 eflow_tt 未找到该工单，规则引擎无法判定"
            return out

        out["found"] = True
        out["level_status"] = _s(row.get("level_status")) or "unknown"
        out["current_severity"] = _s(row.get("ticket_severity"))
        out["expected_severity"] = _s(row.get("expected_severity"))
        rule_id = _s(row.get("matched_rule_id"))
        if not rule_id:
            out["error"] = out["error"] or _s(res.get("error")) or "规则未命中"
            out["available"] = True   # 判定结论有效：确实是「规则未命中」
            return out

        out["matched"] = _matched_block(
            rule_id=rule_id,
            match_details=_s(row.get("match_details")),
            priority=_s(row.get("matched_priority")),
            expected_severity=_s(row.get("expected_severity")),
            changed=out["level_status"] == "adjust",
        )
        out["available"] = True
        # 规则表行数：优先用摘要里的规则数，取不到用当前解析到的行数兜底
        if out["rules_file"] and out["rules_file"].get("rules_count") is None:
            try:
                out["rules_file"]["rules_count"] = len(_load_rule_rows())
            except Exception:  # noqa: BLE001
                pass
        return out
    except Exception as e:  # noqa: BLE001
        out["error"] = out["error"] or f"规则匹配失败：{type(e).__name__}: {e}"
        logger.warning("TT 规则证据匹配失败：%s", out["error"])
        return out
    finally:
        out["elapsed_sec"] = round(time.perf_counter() - t0, 2)


# 判定状态 → 人读文案（与前端 i18n 的 qa.tt.verdict* 一一对应）
_VERDICT_ZH = {"ok": "正常（无需调整）", "adjust": "需调整", "unknown": "无法判定"}
_VERDICT_EN = {"ok": "OK (no change needed)", "adjust": "Needs adjustment", "unknown": "Undetermined"}


def render_evidence_text(evidence: Dict[str, Any], lang: str = "zh") -> str:
    """把证据渲染成给大模型的「本地规则引擎预判定」文本块（zh/en）。

    只要**级别规则或推单规则**任一给出了结论就返回非空串；两者都不可用时返回 ""，
    由调用方保持原有的纯 AI 判定行为。
    （推单规则与级别规则的数据源不同：级别规则表缺失时推单规则仍可能可用，
    所以不能只看 `evidence["available"]` 就整块丢掉。）
    """
    if not evidence or not (evidence.get("available")
                            or (evidence.get("push_rules") or {}).get("available")):
        return ""
    en = lang == "en"
    verdict = (_VERDICT_EN if en else _VERDICT_ZH).get(evidence.get("level_status"), "")
    rf = evidence.get("rules_file") or {}
    xlsx = (rf.get("xlsx") or {})
    lines = [
        "【Local rule engine pre-judgement (authoritative, use this as the source of truth)】"
        if en else "【本地规则引擎预判定（权威口径，与 Email Reports → eFlow TT Rules Check 同源）】"
    ]

    if evidence.get("found"):
        lines.append(
            f"- Current severity: {evidence.get('current_severity') or '-'}"
            if en else f"- 工单当前级别：{evidence.get('current_severity') or '-'}"
        )
        if verdict:
            lines.append(
                f"- Verdict: {verdict}"
                if en else f"- 判定：{verdict}"
            )

    m = evidence.get("matched") or {}
    if m:
        lines.append(
            f"- Matched rule: ID {m.get('rule_id')} | {m.get('rule_name')} | "
            f"priority {m.get('priority')} | target severity {m.get('severity')}"
            if en else
            f"- 命中规则：ID {m.get('rule_id')} | {m.get('rule_name')} | "
            f"优先级 {m.get('priority')} | 目标级别 {m.get('severity')}"
        )
        if m.get("match_details"):
            lines.append(
                f"- Matched conditions: {m['match_details']}"
                if en else f"- 命中条件：{m['match_details']}"
            )
        if m.get("ticket_type") or m.get("trigger_step"):
            lines.append(
                f"- Rule scope: ticket type {m.get('ticket_type') or '-'} / "
                f"step {m.get('trigger_step') or '-'}"
                if en else
                f"- 规则触发：工单类型 {m.get('ticket_type') or '-'} / 触发步骤 {m.get('trigger_step') or '-'}"
            )
        if m.get("creator") or m.get("updater"):
            lines.append(
                f"- Rule owner: created by {m.get('creator') or '-'} / "
                f"updated by {m.get('updater') or '-'} at {m.get('updated_at') or '-'}"
                if en else
                f"- 规则维护人：创建人 {m.get('creator') or '-'} / 更新人 {m.get('updater') or '-'} "
                f"/ 更新于 {m.get('updated_at') or '-'}"
            )
    elif evidence.get("error"):
        lines.append(f"- Rule engine note: {evidence['error']}"
                     if en else f"- 规则引擎说明：{evidence['error']}")

    # 推单规则（TT Push Order Rule）：决定「FO 接单时是否自动建 WO、推给谁、谁自动处理」
    pr = evidence.get("push_rules") or {}
    if pr.get("available"):
        pm = pr.get("matched") or []
        lines.append("【Push order rule / 推单规则（TT Push Order Rule，triggerStep=FOAccept）】"
                     if en else "【推单规则（TT Push Order Rule，触发步骤 FOAccept）】")
        if pm:
            for r in pm[:3]:
                lines.append(
                    f"- Rule Name: {r.get('rule_name')} (ID {r.get('rule_id')}, priority {r.get('priority')})"
                    if en else
                    f"- Rule Name：{r.get('rule_name')}（ID {r.get('rule_id')}，优先级 {r.get('priority')}）"
                )
                lines.append(f"- Rule Type: {r.get('rule_type_text') or r.get('rule_type')}"
                             if en else f"- Rule Type：{r.get('rule_type_text') or r.get('rule_type')}")
                lines.append(f"- Rule Trigger Step: {r.get('trigger_step_text') or r.get('trigger_step')}"
                             if en else f"- Rule Trigger Step：{r.get('trigger_step_text') or r.get('trigger_step')}")
                if r.get("push_path"):
                    lines.append(f"- Push Order Path: {r['push_path']}"
                                 if en else f"- Push Order Path：{r['push_path']}")
                if r.get("auto_operator"):
                    lines.append(f"- *Auto Operator: {r['auto_operator']}"
                                 if en else f"- *Auto Operator：{r['auto_operator']}")
                if r.get("target_handler"):
                    lines.append(f"- Target Handler: {r['target_handler']}"
                                 if en else f"- Target Handler：{r['target_handler']}")
                if r.get("match_details"):
                    lines.append(f"- Matched conditions: {r['match_details']}"
                                 if en else f"- 命中条件：{r['match_details']}")
        else:
            lines.append(
                "- No ACTIVE push order rule matched at FOAccept → the TT is NOT auto-pushed to a WO; "
                "FO must accept and decide manually."
                if en else
                "- FOAccept 步骤**无** ACTIVE 推单规则命中 → 本单不会自动推 WO，需 FO 人工接单后自行判断。"
            )
        if pr.get("other_steps_count"):
            lines.append(
                f"- (Also matched at other steps: {pr['other_steps_count']} rule(s), "
                f"e.g. {', '.join(x.get('rule_name') or '' for x in (pr.get('other_steps') or [])[:3])})"
                if en else
                f"- （另有 {pr['other_steps_count']} 条规则在其他触发步骤命中："
                f"{'、'.join(x.get('rule_name') or '' for x in (pr.get('other_steps') or [])[:3])}）"
            )
        if pr.get("undecidable"):
            lines.append(
                f"- {pr['undecidable']} rule(s) reference fields unavailable locally and were not judged: "
                f"{'; '.join(pr.get('undecidable_rules') or [])}"
                if en else
                f"- 另有 {pr['undecidable']} 条规则的条件含本地取不到的字段，未参与判定："
                f"{'；'.join(pr.get('undecidable_rules') or [])}"
            )
    elif pr.get("error"):
        lines.append(f"- Push order rule note: {pr['error']}"
                     if en else f"- 推单规则说明：{pr['error']}")

    if xlsx.get("name"):
        meta = []
        if rf.get("rules_count") is not None:
            meta.append(f"ACTIVE rules: {rf['rules_count']}" if en else f"ACTIVE 规则 {rf['rules_count']} 条")
        if rf.get("rules_updated_at"):
            meta.append(("file updated " if en else "规则表更新于 ") + rf["rules_updated_at"])
        if rf.get("collected_at"):
            meta.append(("last collect+match " if en else "最近一次收集+匹配 ") + rf["collected_at"])
        lines.append(
            (f"- Rule source file: {xlsx['name']} ({xlsx.get('rel_path') or ''}"
             if en else f"- 规则来源文件：{xlsx['name']}（{xlsx.get('rel_path') or ''}")
            + ("; " + "; ".join(meta) if meta else "") + (")" if en else "）")
        )

    h = evidence.get("handler") or {}
    if h:
        parts = []
        for key, label_zh, label_en in (
            ("last_processer", "最后处理人", "Last handler"),
            ("last_processer_group", "处理组", "Handler group"),
            ("last_process_time", "最后处理时间", "Last handled at"),
            ("organization_id", "组织", "Organization"),
            ("vendor", "厂商", "Vendor"),
        ):
            if h.get(key):
                parts.append(f"{label_en if en else label_zh} {h[key]}")
        if parts:
            # 用 " · " 而不是 " / "：源数据里本身可能含 " /"（实测 last_processer_group
            # 为 `G:3rd_Power;\n /`），同种分隔符会让人误读成多了一个字段。
            lines.append(
                ("- Ticket handling: " if en else "- 工单处理：") + " · ".join(parts)
            )

    lines.append(
        "Requirements: the conclusion MUST name the matched rule (name + ID) and the rule source file, "
        "list the matched condition fields with their actual values, and MUST NOT invent rules outside the "
        "provided rule table."
        if en else
        "要求：结论必须写明「命中规则名称 + 规则ID」与「规则来源文件」，逐条列出命中条件的字段与实际取值；"
        "不得引用规则表之外的规则。"
    )
    return "\n".join(lines)


# ------------------------------------------------------------------ #
# 程序化结论（2026-09-23 作为模型兜底新增；2026-10-19 起是页面唯一结论来源）
# ------------------------------------------------------------------ #

# 结论里要列出的工单字段（用 eflow_tt 的英文列名，与语言无关，取值来自 inspect_tt 报告）
_FACT_KEYS = (
    "site_name", "alarm_occur_time", "created_time", "last_modified_time",
    "alarm_code", "alarm_id", "aid", "alarm_severity", "vip_site", "site_type",
    "down_site_count", "ticket_severity", "status",
    "root_cause", "sub_root_cause",
    "last_processer_group", "last_processer",
)

# 字段 → 表头展示名（zh/en）；未登记的列直接用列名
_FACT_LABEL_EN = {
    "site_name": "Site", "alarm_occur_time": "Alarm time", "created_time": "Ticket created",
    "last_modified_time": "Last modified", "alarm_code": "alarmCode", "alarm_id": "alarm_id",
    "aid": "AID", "alarm_severity": "Alarm severity", "vip_site": "vipSite",
    "site_type": "site_type", "down_site_count": "downSiteCount",
    "ticket_severity": "Current severity", "status": "Status",
    "root_cause": "Root cause", "sub_root_cause": "Sub root cause",
    "last_processer_group": "Handler group", "last_processer": "Last handler",
}
_FACT_LABEL_ZH = {
    "site_name": "站点", "alarm_occur_time": "告警发生时刻", "created_time": "派单时刻",
    "last_modified_time": "最后更新", "alarm_code": "告警码 alarmCode", "alarm_id": "alarm_id",
    "aid": "AID", "alarm_severity": "告警级别", "vip_site": "vipSite",
    "site_type": "site_type", "down_site_count": "downSiteCount",
    "ticket_severity": "当前级别", "status": "状态",
    "root_cause": "根因", "sub_root_cause": "子根因",
    "last_processer_group": "处理组", "last_processer": "最后处理人",
}


def _fact_rows(report: Optional[Dict[str, Any]]) -> list:
    """从 inspect_tt 报告里按 `_FACT_KEYS` 顺序取工单字段（缺失的跳过，不臆造）。

    本地库里的数值列常被读成 float（`alarm_code` / `alarm_id` 会带 `.0`），
    这里统一去掉整数尾巴，与 `inspect_tt.num_tail` 的展示口径一致。
    """
    vals: Dict[str, str] = {}
    for sec in (report or {}).get("fields") or []:
        for it in sec.get("items") or []:
            k = _s(it.get("key"))
            v = _s(it.get("value"))
            if v.endswith(".0") and v[:-2].lstrip("-").isdigit():
                v = v[:-2]
            if k and v and k not in vals:
                vals[k] = v
    return [(k, vals[k]) for k in _FACT_KEYS if vals.get(k)]


def _model_input(report: Optional[Dict[str, Any]]) -> str:
    """取「下发给模型的技能入参」原文（skill_prompt 首行冒号后的字段串）。

    取不到时回退到 `tt_prompt_fields`（裸列名，仅供排查），保证兜底结论里
    始终能看到「模型当时看到的是什么」。
    """
    sp = _s((report or {}).get("skill_prompt"))
    first = sp.splitlines()[0].strip() if sp else ""
    if "：" in first:
        tail = first.split("：", 1)[1].strip()
        if "=" in tail:
            return tail
    tpf = (report or {}).get("tt_prompt_fields") or {}
    return " ".join(f"{k}={_s(v)}" for k, v in tpf.items() if _s(v))


def render_local_conclusion(evidence: Dict[str, Any],
                            report: Optional[Dict[str, Any]] = None,
                            lang: str = "zh",
                            reason: str = "") -> str:
    """**程序化**生成 TT 判定结论（Markdown）。

    用途有两处：`GET /api/eflow-qa/tt-ai-fast` 的**主结论**（页面「TT 快速分析」
    2026-10-19 起只走这条，不再调用模型），以及 `GET /api/eflow-qa/tt-ai`
    在候选模型档全失败时的兜底结论。

    这里没有任何模型参与，全部来自本地规则引擎与工单字段，
    因此结论里的规则名/级别/条件与 Email Reports → eFlow TT Rules Check 页签逐字一致。

    规则明细的完整版（逐条匹配、下载链接）由页面的「规则命中证据」块展示，
    本函数只给一行摘要，避免同一份内容在页面上出现两次。

    返回 Markdown 文本；`reason` 为「为什么没用模型结论」（模型失败原因或本地化开关），
    非空时才附在末尾，便于复制/归档时自证；程序化主结论应传空串。
    """
    en = lang == "en"
    ev = evidence or {}
    m = ev.get("matched") or {}
    rf = ev.get("rules_file") or {}
    xlsx = rf.get("xlsx") or {}
    h = ev.get("handler") or {}
    label = _FACT_LABEL_EN if en else _FACT_LABEL_ZH

    out: list = []
    out.append("## TT severity verdict (programmatic / rule engine)"
               if en else "## TT 工单级别判定（程序化 · 本地规则引擎）")
    out.append("")
    out.append(
        "> This verdict is computed directly by the local rule engine from the eFlow "
        "*TT Update Severity Rule* table — the same source as the daily 04:20 "
        "`eflow_tt_rules_check` job. No model is involved."
        if en else
        "> 本结论由**本地规则引擎**按 eFlow「TT Update Severity Rule」直接算出，"
        "与每天 04:20 的 `eflow_tt_rules_check` 定时任务同源，**无模型参与**。"
    )
    out.append("")

    # 1) 工单关键字段
    facts = _fact_rows(report)
    rule_input = _model_input(report)
    if facts or rule_input:
        out.append("### 1. Ticket fields / 工单关键字段" if en else "### 1. 工单关键字段")
        out.append("")
        out.append("| Field | Value |" if en else "| 字段 | 取值 |")
        out.append("| --- | --- |")
        for k, v in facts:
            out.append(f"| {label.get(k, k)} | {v} |")
        out.append("")
        if rule_input:
            out.append(("- Alarm-side fields (skill prompt text): `" if en
                        else "- 告警侧字段（技能 prompt 原文）：`")
                       + rule_input + "`")
            out.append("")

    # 2) 判定结论
    out.append("### 2. Verdict / 判定结论" if en else "### 2. 判定结论")
    out.append("")
    if ev.get("available") and ev.get("found"):
        cur = ev.get("current_severity") or "-"
        exp = ev.get("expected_severity") or "-"
        status = _s(ev.get("level_status"))
        verdict = (_VERDICT_EN if en else _VERDICT_ZH).get(status, "-")
        if m:
            out.append(
                f"- Matched rule: **{m.get('rule_name')}** (ID {m.get('rule_id')}, "
                f"priority {m.get('priority')}) → target severity **{m.get('severity')}**"
                if en else
                f"- 命中规则：**{m.get('rule_name')}**（ID {m.get('rule_id')}，"
                f"优先级 {m.get('priority')}）→ 目标级别 **{m.get('severity')}**"
            )
            if m.get("match_details"):
                out.append(("- Condition check: " if en else "- 逐条匹配：") + str(m["match_details"]))
        out.append(f"- Current severity: **{cur}** / Rule severity: **{exp}**" if en
                   else f"- 当前级别：**{cur}**　规则建议级别：**{exp}**")
        out.append(f"- **Verdict: {verdict}**" if en else f"- **判定：{verdict}**")
        out.append("")
        out.append(
            "> Note: the rule engine evaluates the `eflow_tt` columns (the values in the condition "
            "check above are the real inputs); the alarm-side fields above may come from the alarm row "
            "and differ on fields like `downSiteCount`. Trust the condition check."
            if en else
            "> 说明：规则引擎按 `eflow_tt` 表字段判定（上面「逐条匹配」里的就是实际取值），"
            "与「告警侧字段」不一定同源（如 `downSiteCount` 可能取自告警行），**以逐条匹配为准**。"
        )
    elif ev.get("available") and not ev.get("found"):
        out.append("- The ticket is not in the local `eflow_tt` mirror, so the rule engine cannot "
                   "match it; use the eFlow live record to confirm the severity."
                   if en else
                   "- 本地 `eflow_tt` 无该工单（仅存在于 eFlow 实时接口），规则引擎无法匹配；"
                   "请以 eFlow 侧工单记录核对级别。")
    else:
        out.append("- Rule engine unavailable: " + (_s(ev.get("error")) or "unknown reason")
                   if en else "- 规则引擎不可用：" + (_s(ev.get("error")) or "原因未知"))
        out.append("- Severity verdict: **undetermined** — needs manual confirmation."
                   if en else "- 级别判定：**无法判定**，需人工确认。")
    out.append("")

    # 3) 规则来源与责任人
    pr = ev.get("push_rules") or {}
    push_hit = pr.get("matched") or []
    if xlsx.get("name") or h or push_hit or pr.get("available"):
        out.append("### 3. Rule source & owners / 规则来源与责任人" if en
                   else "### 3. 规则来源与责任人")
        out.append("")
        if xlsx.get("name"):
            meta = []
            if rf.get("rules_count") is not None:
                meta.append(f"ACTIVE rules: {rf['rules_count']}" if en
                            else f"ACTIVE 规则 {rf['rules_count']} 条")
            if rf.get("rules_updated_at"):
                meta.append(("file updated " if en else "规则表更新于 ") + rf["rules_updated_at"])
            if rf.get("collected_at"):
                meta.append(("last collect+match " if en else "最近一次收集+匹配 ") + rf["collected_at"])
            bits = [b for b in (xlsx.get("rel_path") or "", *meta) if b]
            if en:
                out.append(f"- Rule source file: `{xlsx['name']}` ({'; '.join(bits)})")
            else:
                out.append(f"- 规则来源文件：`{xlsx['name']}`（{'；'.join(bits)}）")
        if m and (m.get("creator") or m.get("updater")):
            out.append(
                f"- Rule owner: created by {m.get('creator') or '-'} / updated by "
                f"{m.get('updater') or '-'} at {m.get('updated_at') or '-'}"
                if en else
                f"- 规则维护人：创建人 {m.get('creator') or '-'} / 更新人 {m.get('updater') or '-'}"
                f"（更新于 {m.get('updated_at') or '-'}）"
            )
        if h:
            parts = [str(h[k]) for k in ("last_processer", "last_processer_group",
                                         "last_process_time", "organization_id", "vendor") if h.get(k)]
            if parts:
                out.append(("- Ticket handler: " if en else "- 工单处理人：") + " · ".join(parts))
        out.append("")

        # 3b) 推单规则：谁推单、推给谁、谁自动处理（字段名沿用 eFlow 规则页签写法）
        prf = pr.get("rules_file") or {}
        out.append("#### 推单规则（TT Push Order Rule，触发步骤 FOAccept）" if not en
                   else "#### Push order rule (TT Push Order Rule, triggerStep=FOAccept)")
        out.append("")
        if push_hit:
            for r in push_hit[:3]:
                out.append(f"- Rule Name: {r.get('rule_name') or '-'}"
                           + (f"（ID {r.get('rule_id')}，优先级 {r.get('priority')}）" if not en
                              else f" (ID {r.get('rule_id')}, priority {r.get('priority')})"))
                out.append(f"- Rule Type: {r.get('rule_type_text') or r.get('rule_type') or '-'}")
                out.append(f"- Rule Trigger Step: {r.get('trigger_step_text') or r.get('trigger_step') or '-'}")
                out.append(f"- Push Order Path: {r.get('push_path') or '-'}")
                out.append(f"- *Auto Operator: {r.get('auto_operator') or '-'}")
                out.append(f"- Target Handler: {r.get('target_handler') or '-'}")
                if r.get("match_details"):
                    out.append(("- Matched conditions: " if en else "- 命中条件：") + str(r["match_details"]))
            out.append("")
            if pr.get("other_steps_count"):
                names = "、".join(x.get("rule_name") or "" for x in (pr.get("other_steps") or [])[:3])
                out.append(("- Other steps also matched: " if en else "- 其他触发步骤另有命中：")
                           + (f"{pr['other_steps_count']} rule(s) ({names})" if en
                              else f"{pr['other_steps_count']} 条（{names}）"))
                out.append("")
        elif pr.get("available"):
            out.append("- FOAccept 步骤**无** ACTIVE 推单规则命中 → 本单不会自动推 WO，"
                       "由 FO 人工接单后自行判断是否建单。" if not en else
                       "- No ACTIVE push order rule matched at FOAccept → the TT is not auto-pushed "
                       "to a WO; FO accepts and decides manually.")
            out.append("")
        if pr.get("available") and prf.get("name"):
            meta = []
            if prf.get("rules_count") is not None:
                meta.append(f"ACTIVE rules: {prf['rules_count']}" if en
                            else f"ACTIVE 推单规则 {prf['rules_count']} 条")
            if prf.get("mtime_text"):
                meta.append(("file updated " if en else "规则表更新于 ") + prf["mtime_text"])
            if pr.get("undecidable"):
                meta.append((f"{pr['undecidable']} rule(s) not judged (fields unavailable locally)"
                             if en else f"{pr['undecidable']} 条条件字段本地不可评估，未参与判定"))
            bits = [b for b in (prf.get("rel_path") or "", *meta) if b]
            out.append((f"- Rule source file: `{prf['name']}` ({'; '.join(bits)})" if en
                        else f"- 推单规则来源文件：`{prf['name']}`（{'；'.join(bits)}）"))
            out.append("")
        elif pr.get("error"):
            out.append(("- Push order rule note: " if en else "- 推单规则说明：") + str(pr["error"]))
            out.append("")

    # 4) 处置建议（按判定分支，全部可程序化推出）
    status = _s(ev.get("level_status"))
    out.append("### 4. Recommended actions / 处置建议" if en else "### 4. 处置建议")
    out.append("")
    if ev.get("available") and ev.get("found") and status == "adjust":
        out.append(
            f"1. Adjust the ticket severity from **{ev.get('current_severity') or '-'}** to "
            f"**{ev.get('expected_severity') or '-'}** per rule ID {m.get('rule_id') or '-'}, and re-check "
            "whether the dispatch group must change with the severity."
            if en else
            f"1. 按命中规则 ID {m.get('rule_id') or '-'} 将工单级别由 "
            f"**{ev.get('current_severity') or '-'}** 调整为 **{ev.get('expected_severity') or '-'}**，"
            "并同步核对派单组是否随级别变化。"
        )
        out.append("2. Record the adjustment reason (rule name + ID) in the ticket history."
                   if en else "2. 在工单备注中记录调整依据（规则名称 + 规则ID），便于复盘追溯。")
    elif ev.get("available") and ev.get("found") and status == "ok":
        out.append(
            f"1. Keep the current severity **{ev.get('current_severity') or '-'}** and proceed with the "
            "normal handling flow."
            if en else
            f"1. 维持当前级别 **{ev.get('current_severity') or '-'}**，按既有流程处理。"
        )
        out.append(
            "2. Re-run the rule check if `alarmCode`, `vipSite` or `assetType` changes, or if the alarm "
            "escalates — the matched rule may then differ."
            if en else
            "2. 若 `alarmCode`、`vipSite`、`assetType` 发生变化或告警升级，需重新套用规则 —— "
            "命中规则可能随之改变。"
        )
    elif ev.get("available") and not ev.get("found"):
        out.append("1. Verify the ticket in eFlow and confirm its severity manually."
                   if en else "1. 到 eFlow 侧核对该工单并人工确认级别。")
        out.append("2. If the ticket should be visible locally, check the eflow_sync jobs."
                   if en else "2. 若该单本应在本地可见，检查 eflow_sync 系列同步任务。")
    else:
        out.append("1. Confirm the severity manually; the rule engine could not produce a verdict."
                   if en else "1. 规则引擎未能给出判定，请人工确认级别。")
        out.append("2. Escalate to the rule architect if the rule table is missing or unreadable."
                   if en else "2. 若规则表缺失或不可读，请联系规则架构师处理。")
    out.append("")

    if reason:
        out.append("---")
        out.append("")
        out.append(f"*未采用模型结论的原因：`{reason}`*" if not en
                   else f"*Reason no model verdict was used: `{reason}`*")
    return "\n".join(out)


__all__ = ["build_rule_evidence", "build_push_rule_evidence",
           "render_evidence_text", "render_local_conclusion"]
