# -*- coding: utf-8 -*-
"""
MS SLA / NUR 服务集成
整合 ran_ume_sla.py + update_contract_kpis_report.py
功能：
  1. RAN UME SLA 指标同步（API 查询 → xlsx → SQLite → API写入远程/FTP回退）
  2. Contract KPIs 报表自动更新（FTP 下载 → 数据更新 → 导出 → FTP 上传）
"""
import base64
import gc
import glob
import gzip
import io
import json
import os
import re
import shutil
import socket
import sqlite3
import ssl
import subprocess
import sys
import time
import copy
import logging
import warnings
import threading
import random
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import contextmanager
from datetime import datetime, date, timezone, timedelta
from ftplib import FTP, error_perm
from pathlib import Path

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
from backend.database import get_raw_connection

try:
    from .atomic_write import atomic_save_workbook, AtomicReplaceFallback, cleanup_stale_tmp_files
except ImportError:
    logging.getLogger(__name__).warning("[SILENT-EXCEPT] <module:ms_sla_service>: exception swallowed", exc_info=True)
    from atomic_write import atomic_save_workbook, AtomicReplaceFallback, cleanup_stale_tmp_files

try:
    from PIL import Image, ImageDraw, ImageFont
    HAS_PIL = True
except ImportError:
    logging.getLogger(__name__).warning("[SILENT-EXCEPT] <module:ms_sla_service>: exception swallowed", exc_info=True)
    HAS_PIL = False

_SUBPROCESS_FLAGS = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# 远程写入 API 配置（替代 FTP 上传整个 db 文件）
try:
    from backend.services.db_writer_client import DbWriterClient
except ImportError:
    logging.getLogger(__name__).warning("[SILENT-EXCEPT] <module:ms_sla_service>: exception swallowed", exc_info=True)
    from db_writer_client import DbWriterClient
try:
    _db_writer = DbWriterClient()
except Exception:
    _db_writer = None
    print("[WARNING] db_writer_client 初始化失败，将回退到 FTP 模式")

logger = logging.getLogger("ms_sla_service")

# ============================================================
# 路径配置
# ============================================================
DATA_DIR = Path(r"d:\01_RPA\Task\my_app\data\ms_sla")
DATA_DIR.mkdir(parents=True, exist_ok=True)

DB_NAME = "lbr_ms_sla_kpi.sqlite3"
DB_PATH = DATA_DIR / DB_NAME

# 报表文件
TARGET_FILE = DATA_DIR / "ZTE OLR MS Contract KPIs & WOs Achievement Auto-Reports.xlsx"
KPI_LIST_FILE = DATA_DIR / "KPI list - Operation.xlsx"

# SLA 导出临时目录
SLA_OUTPUT_DIR = str(DATA_DIR / "sla_exports")
os.makedirs(SLA_OUTPUT_DIR, exist_ok=True)

# ============================================================
# FTP 配置
# ============================================================
# 上传 FTP（ZTEEMS 账号）
# FTP_UP_HOST 已收敛到 backend.config.remote_host.get_ftp_host()（见 ms_sla_sync / ms_sla_report）
FTP_UP_PORT = 21
FTP_UP_USER = "ZTEEMS"
FTP_UP_PASS = "DX"
FTP_UP_REMOTE_DIR = "/01_RPA/LBR_Tools/KPI/MS_SLA"

# 下载 FTP（FINEBI 账号）
# FTP_SRC_HOST 已收敛到 backend.config.remote_host.get_ftp_host()
FTP_SRC_PORT = 21
FTP_SRC_USER = "FINEBI"
FTP_SRC_PASS = "DX"
FTP_SRC_BASE_DIR = "/webroot/WEB-INF/schedule/NUR_SLA"

# SQLite 批量插入大小
BATCH_SIZE = 2000

# ============================================================
# RAN UME API 配置
# ============================================================
UME_HOST = "10.100.28.27"
UME_PORT = 28001
# 只留用户名（非敏感）；口令不得写在这里，由 credential_store 的 `zte_ume` 组提供
# （取值点见 ms_sla_sync._runtime_credentials）。2026-09-27 删除硬编码 UME_PASSWORD。
UME_USERNAME = "NOC"


# ============================================================
# 外部化配置（Settings →「任务配置」维护的 ms_sla_config.json）
#   优先级：环境变量 > 配置文件 > 内置默认；FTP/UME 账号密码不进配置文件（走「凭据管理」）。
#   ⚠ 必须位于 ms_sla_sync / ms_sla_report 里 `from ms_sla_service import …` 取值之前：
#     那两个模块是值拷贝，覆盖晚了就只影响本模块自身。
# ============================================================
try:
    from backend.services.task_config_files import apply_config_overrides
except ImportError:
    from task_config_files import apply_config_overrides  # type: ignore

apply_config_overrides(
    globals(),
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "ms_sla_config.json"),
    {
        "ume_host": ("UME_HOST", str, "MS_SLA_UME_HOST"),
        "ume_port": ("UME_PORT", int, "MS_SLA_UME_PORT"),
        "batch_size": ("BATCH_SIZE", lambda v: max(1, int(v)), None),
        "ftp_up_port": ("FTP_UP_PORT", int, None),
        "ftp_up_remote_dir": ("FTP_UP_REMOTE_DIR", str, "MS_SLA_FTP_DIR"),
        "ftp_src_port": ("FTP_SRC_PORT", int, None),
        "ftp_src_base_dir": ("FTP_SRC_BASE_DIR", str, "MS_SLA_SRC_DIR"),
    },
    log_name="ms_sla",
)

# ============================================================
# 统一同步入口（SLA 同步 + Contract KPIs 报表更新）
# 注：数据同步实现已拆至 ms_sla_sync.py，报表生成已拆至 ms_sla_report.py（Phase 4）
# ============================================================

def run_full_sync():
    """完整同步流程：先执行 SLA 指标同步，再执行 Contract KPIs 报表更新。"""
    # Phase 4 拆分：实现下沉子模块，函数内显式导入（模块内全局名不走 __getattr__）
    from backend.services.ms_sla_sync import run_ms_sla_sync
    from backend.services.ms_sla_report import run_contract_kpis_report

    total_start = time.time()
    logger.info("=" * 60)
    logger.info("完整同步开始: SLA 同步 + Contract KPIs 报表更新")
    logger.info("=" * 60)

    # 清理上次异常退出残留的临时文件
    cleanup_stale_tmp_files(DATA_DIR)

    result = {"success": True}

    # Step 1: SLA 同步
    logger.info("\n----- Step 1/2: RAN UME SLA 同步 -----")
    sla_stats = run_ms_sla_sync()
    result["sla_sync"] = sla_stats
    if not sla_stats.get("success"):
        logger.warning("SLA 同步失败，继续执行报表更新...")
        result["success"] = False

    # Step 2: Contract KPIs 报表更新
    logger.info("\n----- Step 2/2: Contract KPIs 报表更新 -----")
    report_stats = run_contract_kpis_report()
    result["report"] = report_stats
    if not report_stats.get("success"):
        logger.warning("Contract KPIs 报表更新失败")
        result["success"] = False

    elapsed = time.time() - total_start
    logger.info("=" * 60)
    logger.info(f"完整同步结束: {elapsed:.1f}s")
    logger.info("=" * 60)
    return result


# ============================================================
# 入口
# ============================================================
if __name__ == "__main__":
    import sys
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        logging.getLogger(__name__).warning("[SILENT-EXCEPT] <module:ms_sla_service>: exception swallowed", exc_info=True)
        pass
    total_start = time.time()
    print(f"[MAIN] MS SLA service started at {time.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"[MAIN] Data dir: {DATA_DIR}")
    print(f"[MAIN] DB: {DB_PATH}")

    mode = sys.argv[1] if len(sys.argv) > 1 else "full"

    # Phase 4 拆分：入口函数已迁移至子模块（模块内全局名不走 __getattr__，需显式导入）
    try:
        from backend.services.ms_sla_sync import run_ms_sla_sync
        from backend.services.ms_sla_report import run_contract_kpis_report
    except ImportError:
        from ms_sla_sync import run_ms_sla_sync
        from ms_sla_report import run_contract_kpis_report

    if mode == "full":
        stats = run_full_sync()
        print(f"\n[MAIN] Full sync completed in {time.time() - total_start:.1f}s")

    elif mode == "sync":
        stats = run_ms_sla_sync()
        print(f"\n[MAIN] SLA sync completed:")
        if stats.get("success"):
            print(f"  Downloaded: {stats['downloaded']} files")
            print(f"  Imported: {stats['total_rows']} rows")
            print(f"  Tables: {stats['tables']}")
            print(f"  FTP: {'OK' if stats.get('ftp_uploaded') else 'FAILED'}")
        else:
            print(f"  Error: {stats.get('error')}")

    elif mode == "report":
        stats = run_contract_kpis_report()
        print(f"\n[MAIN] Contract KPIs report completed:")
        if stats.get("success"):
            print(f"  Daily rows: {stats['daily_rows_updated']}")
            print(f"  Exported: {stats['exported_files']} files")
        else:
            print(f"  Error: {stats.get('error')}")

    elif mode == "help":
        print("Usage: python ms_sla_service.py [mode]")
        print("  modes:")
        print("    full     Complete sync: SLA sync + Contract KPIs report (default)")
        print("    sync     Run RAN UME SLA sync only")
        print("    report   Run Contract KPIs report only")

    else:
        print(f"[MAIN] Unknown mode: {mode}, use 'full', 'sync', or 'report'")

    total_elapsed = time.time() - total_start
    print(f"\n[MAIN] Total elapsed: {total_elapsed:.2f} seconds")


# ========== Phase 4 拆分兼容层（PEP 562 惰性重导出） ==========
# 以下入口函数已迁移至子模块；旧导入路径（from <本模块> import X）通过
# 模块级 __getattr__ 惰性解析，避免顶层互相导入造成循环依赖。
_PHASE4_SPLITS = {
    "run_ms_sla_sync": "ms_sla_sync",
    "run_contract_kpis_report": "ms_sla_report",
    "write_to_db": "ms_sla_sync",
    "read_nur_daily": "ms_sla_report",
    "read_sla_daily": "ms_sla_report",
    "update_wo_sla_report": "ms_sla_report",
}


def __getattr__(name):
    target = _PHASE4_SPLITS.get(name)
    if target is not None:
        import importlib
        pkg = f"{__package__}.{target}" if __package__ else target
        return getattr(importlib.import_module(pkg), name)
    raise AttributeError(f"module {__name__!r} has no attribute {name!r}")
