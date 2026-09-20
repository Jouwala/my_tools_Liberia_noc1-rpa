"""MS Contract KPIs / WOs SLA 报表（MS_SLA_REPORT_WEB_V1）数据服务。

数据源
------
``data/ms_sla`` 下由任务产出的两份**报表工作簿本体**（页面不做二次加工，与邮件
附件是同一份文件）：

    ZTE OLR MS WOs SLA Achievements Daily Report_<YYYY-MM-DD>.xlsx   -> kind="wo_sla"
    ZTE OLR MS Contract KPIs & WOs Achievement Auto-Reports.xlsx      -> kind="contract_kpi"

- ``wo_sla``：任务 ``kpi_ms_wo_sla``（MS WOs SLA 日报邮件）读的日报，单 sheet
  ``WOs SLA achievement Report``；
- ``contract_kpi``：任务 ``ms_sla_full``（MS SLA 完整同步，每天 05:00）就地更新的
  自动化报表工作簿，含 4 张 sheet —— ``Contract KPIs Report (New)`` /
  ``WOs SLA achievement Report`` / ``FineBI Dashboard Planing`` / ``利比里亚自动化报表``。
  其中 ``WOs SLA achievement Report`` 与 ``wo_sla`` 那份**逐字同构**，因此两个页面
  共用同一套解析与渲染。

⚠️ 这两份工作簿的表头**都不在第 1 行**，而且是多层合并表头：

    WOs SLA achievement Report
        行1  标题
        行2  分组头（Date | TTR (%) | TTR_P1_4h (%) | … 5 组）
        行3  目标值（0.98 / 0.98 / 0.95 / 0.9 / 0.9，横向合并到整组）
        行4  叶子列名（Total | Within SLA | Achieve ×5）
        行5  Month 汇总行
        行6+ 日数据

    Contract KPIs Report (New)
        行1  标题
        行2  KPIs Category（Operation KPIs | Network KPIs | PerSite KPIs）
        行3  分组头（Network Unavailability Ratio | Time-to-Restore(TTR) | 2G Voice | …）
        行4  27 个 KPI 列名
        行5  目标值
        行6  Month 汇总行
        行7+ 日数据

所以读取必须按 ``SHEET_LAYOUT`` 的行偏移 + 横向合并前向填充（``_ffill``）还原分组，
否则会把目标行/说明行当数据行。

口径
----
- **指标（metric）粒度**：叶子列名若是 ``Total`` / ``Within SLA`` / ``Achieve``
  这类「角色列」，则该组（如 ``TTR (%)``）整体算一个指标，取值列用 ``Achieve``；
  否则每个叶子列各自是一个指标（Contract KPIs 的 27 个 KPI 即此情形）。
- **目标方向**：KPI 名含 NUR / DCR / Blocking / Congestion / Unavailability 的
  「越小越好」，其余「越大越好」；达标 = 月值 vs 目标值按方向比较。
- 页面所有数字都直接来自工作簿，本服务只做结构化与聚合（min/max/avg/达标天数）。
"""
from __future__ import annotations

import json
import logging
import re
import threading
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)

# .../Task/my_app
MY_APP_DIR = Path(__file__).resolve().parent.parent.parent
MS_SLA_DIR = MY_APP_DIR / "data" / "ms_sla"

#: 支持的报表种类。``pattern`` 精确匹配生成端写出的文件名（HTML/PNG 同名文件不匹配）。
REPORT_KINDS: Dict[str, Dict[str, Any]] = {
    "wo_sla": {
        "label_en": "MS WOs SLA Achievements Daily Report",
        "label_zh": "MS WOs SLA 日报",
        "pattern": r"^ZTE OLR MS WOs SLA Achievements Daily Report_\d{4}-\d{2}-\d{2}\.xlsx$",
        "sheets": ("WOs SLA achievement Report",),
        "task": "kpi_ms_wo_sla",
    },
    "contract_kpi": {
        "label_en": "MS Contract KPIs & WOs Achievement Auto-Reports",
        "label_zh": "MS Contract KPIs 自动化报表",
        "pattern": r"^ZTE OLR MS Contract KPIs & WOs Achievement Auto-Reports\.xlsx$",
        "sheets": (
            "Contract KPIs Report (New)",
            "WOs SLA achievement Report",
            "FineBI Dashboard Planing",
            "利比里亚自动化报表",
        ),
        "task": "ms_sla_full",
    },
}
DEFAULT_KIND = "wo_sla"

#: 叶子列名里的「角色」列 —— 出现这些名字说明整组是一个指标（Total/Within SLA/Achieve）
ROLE_LABELS: Dict[str, str] = {
    "total": "total",
    "within sla": "within",
    "achieve": "achieve",
}

#: 各 sheet 的读取配置（行号均 0 基）
#:   layout          kpi_matrix = 多层表头 + 目标行 + Month 行；flat = 单行表头
#:   category_row    最上层分类行（KPIs Category / Operation KPIs / Network KPIs …）
#:   group_row       分组行
#:   column_row      叶子列名行
#:   target_row      目标值行
#:   data_start      第一条数据行
#:   header_row      flat 的表头行
#:   start_col       flat 的起始列（FineBI 表从 B 列开始）
#:   label_header    矩阵首列的列名（Excel 里由合并单元格提供，需补回）
#:   vertical_fill   需要纵向前向填充的列（合并单元格只保留左上角值）
#:   title_row       标题行（作为 notes 透传）
SHEET_LAYOUT: Dict[str, Dict[str, Any]] = {
    "WOs SLA achievement Report": {
        "layout": "kpi_matrix",
        "group_row": 1,
        "column_row": 3,
        "target_row": 2,
        "data_start": 4,
        "label_header": "Date",
        "title_row": 0,
    },
    "Contract KPIs Report (New)": {
        "layout": "kpi_matrix",
        "category_row": 1,
        "group_row": 2,
        "column_row": 3,
        "target_row": 4,
        "data_start": 5,
        "label_header": "Date",
        "title_row": 0,
    },
    "FineBI Dashboard Planing": {
        "layout": "flat",
        "header_row": 1,
        "start_col": 1,
        "data_start": 2,
        "vertical_fill": ("Domain",),
        "title_row": 0,
        "title_col": 1,
    },
    "利比里亚自动化报表": {
        "layout": "flat",
        "header_row": 0,
        "start_col": 0,
        "data_start": 1,
    },
}

#: 各 sheet 的筛选下拉候选列（flat 用列名；矩阵 sheet 不支持下拉筛选）
FILTER_COLUMNS: Dict[str, Tuple[Tuple[str, str], ...]] = {
    "FineBI Dashboard Planing": (
        ("Domain", "domain"),
        ("Priority", "priority"),
    ),
    "利比里亚自动化报表": (
        ("优先级", "priority"),
        ("频率", "frequency"),
        ("团队", "team"),
        ("Data Source", "dataSource"),
        ("Type", "type"),
        ("业务自动化类型", "autoType"),
        ("推动方式", "pushMode"),
    ),
}

#: sheet 在页面上的 i18n 文案键（前端用；这里也返回一份便于排障）
SHEET_LABEL_KEY: Dict[str, str] = {
    "Contract KPIs Report (New)": "msKpi.sheet.contract",
    "WOs SLA achievement Report": "msKpi.sheet.woSla",
    "FineBI Dashboard Planing": "msKpi.sheet.finebi",
    "利比里亚自动化报表": "msKpi.sheet.automation",
}

#: 「越小越好」的 KPI 关键字（目标方向判定）
_LOWER_KEYWORDS = ("nur", "dcr", "blocking", "congestion", "unavailability")

#: PerSite KPIs 的特殊口径：这一组的值不是站点级原始指标，而是「达标站点占比」，
#: 目标统一 90% —— sheet 里的分组名就叫 "2G (90%)" / "3G (90%)"，与生成端
#: ``ms_sla_report.apply_conditional_formatting()`` 里的 ``v >= 0.9`` 判定完全一致。
#: 目标行里 PerSite 那一段写的是站点级原始门限（350 / 0.02 / 0.97 …），直接拿来
#: 跟占比比较会把「98.7% 站点达标」判成未达标，所以这里必须覆盖。
PERSITE_CATEGORY = "PerSite KPIs"
PERSITE_TARGET = 0.9

DEFAULT_PAGE_SIZE = 50
MAX_PAGE_SIZE = 5000


def normalize_kind(kind: Optional[str]) -> str:
    """非法/缺失的 kind 一律回落到默认值（页面不该因为一个参数 500）。"""
    k = (kind or "").strip().lower()
    return k if k in REPORT_KINDS else DEFAULT_KIND


# =============================================================================
# 通用小工具
# =============================================================================
def _to_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        # 这两份报表里大量使用不换行空格（\xa0）做视觉缩进（"FO "、"CS\xa0"），
        # 统一换成普通空格再 strip，否则前端会显示出看不见的尾随空白。
        return v.replace("\xa0", " ").strip()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return _period_label(v)
    return str(v).strip()


def _period_label(v: Any) -> str:
    """日期单元格 -> ``YYYY-MM-DD``（Excel 里是 2026-09-01 00:00:00 这样的 datetime）。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return _to_text(v)


def _to_num(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("%", "")
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _json_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return _period_label(v)
    if isinstance(v, str):
        s = v.replace("\xa0", " ").strip()
        return s if s else None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _ffill(values: Sequence[Any]) -> List[Any]:
    """横向前向填充 —— 还原 Excel 横向合并单元格（只有左上角有值）。"""
    out: List[Any] = []
    last: Any = None
    for v in values:
        if v is not None and v != "":
            last = v
        out.append(last)
    return out


def _round(v: Optional[float], nd: int = 6) -> Optional[float]:
    return None if v is None else round(v, nd)


def _lower_is_better(*labels: str) -> bool:
    text = " ".join(x for x in labels if x).lower()
    return any(k in text for k in _LOWER_KEYWORDS)


def _met(value: Optional[float], target: Optional[float], lower: bool) -> Optional[bool]:
    if value is None or target is None:
        return None
    return bool(value <= target) if lower else bool(value >= target)


def _col_format(label: str, category: str, role: str) -> str:
    """列/指标的展示口径：``percent``（0~1 的比率，前端 ×100 显示）或 ``number``。

    - ``Achieve`` 角色列一定是达成率；
    - 列名带 ``(%)`` 的是百分比指标；
    - PerSite 那一段是「达标站点占比」，也按百分比显示。
    """
    if role == "achieve" or "%" in label or category == PERSITE_CATEGORY:
        return "percent"
    return "number"


def _normalize_columns(raw: Iterable[Any]) -> List[str]:
    out: List[str] = []
    seen: Dict[str, int] = {}
    for i, c in enumerate(raw):
        name = _to_text(c) or f"Column {i + 1}"
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name}#{seen[name]}")
    return out


# =============================================================================
# 工作簿读取与缓存
# =============================================================================
def latest_file(kind: str = DEFAULT_KIND) -> Optional[Path]:
    """某类报表的工作簿（wo_sla 取时间戳最新的那份）。"""
    kind = normalize_kind(kind)
    if not MS_SLA_DIR.is_dir():
        return None
    pattern = re.compile(REPORT_KINDS[kind]["pattern"])
    hits = [p for p in MS_SLA_DIR.iterdir()
            if p.is_file() and pattern.match(p.name)]
    if not hits:
        return None
    return max(hits, key=lambda p: p.stat().st_mtime)


def _cache_key(paths: Sequence[Path]) -> Tuple:
    out: List[Any] = []
    for p in paths:
        try:
            st = p.stat()
            out.append((str(p), int(st.st_mtime), int(st.st_size)))
        except OSError:
            out.append((str(p), None, None))
    return tuple(out)


def _read_sheets(path: Path, names: Sequence[str]) -> Dict[str, List[List[Any]]]:
    """读若干 sheet 的原始单元格（值模式，非公式）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: Dict[str, List[List[Any]]] = {}
    try:
        for name in names:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            # 防御「未声明 dimension → read_only 把整表读成 1 列」（同 sp_report_service）
            reset = getattr(ws, "reset_dimensions", None)
            if callable(reset):
                reset()
            out[name] = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return out


# =============================================================================
# 解析：多层表头矩阵（WOs SLA / Contract KPIs）
# =============================================================================
def _cell(raw: List[List[Any]], r: int, c: int) -> Any:
    if r < 0 or r >= len(raw):
        return None
    row = raw[r]
    return row[c] if c < len(row) else None


def _group_spans(filled: Sequence[str]) -> List[Tuple[str, int, int]]:
    """把（已前向填充的）分组行压成 [(组名, 起列, 止列)]。"""
    spans: List[Tuple[str, int, int]] = []
    start = 1  # 跳过首列（行标签列）
    cur = filled[1] if len(filled) > 1 else ""
    for i in range(2, len(filled)):
        if filled[i] != cur:
            spans.append((cur, start, i - 1))
            cur = filled[i]
            start = i
    if len(filled) > 1:
        spans.append((cur, start, len(filled) - 1))
    return spans


def _parse_kpi_matrix(raw: List[List[Any]], cfg: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """多层合并表头 + 目标行 + Month 行的 KPI 矩阵 -> 结构化矩阵。"""
    ncols = max((len(r) for r in raw), default=0)
    if ncols <= 1:
        return None

    label_header = _to_text(cfg.get("label_header")) or "Date"
    col_row = int(cfg["column_row"])
    group_row = int(cfg["group_row"])
    cat_row = cfg.get("category_row")
    target_row = int(cfg["target_row"])

    leaves = [_to_text(_cell(raw, col_row, i)) for i in range(ncols)]
    groups = _ffill([_to_text(_cell(raw, group_row, i)) for i in range(ncols)])
    cats = (_ffill([_to_text(_cell(raw, cat_row, i)) for i in range(ncols)])
            if cat_row is not None else [""] * ncols)

    # ---- 列模型（按列下标做 key，因为叶子名会重复：如 3G DCR (%) 出现在 3G Voice 与 3G (90%) 下）----
    columns: List[Dict[str, Any]] = []
    for i in range(1, ncols):
        leaf = leaves[i]
        group = _to_text(groups[i])
        cat = _to_text(cats[i])
        if not leaf and not group:
            continue
        label = leaf or group
        role = ROLE_LABELS.get(label.lower(), "value")
        is_persite = cat == PERSITE_CATEGORY
        lower = False if is_persite else _lower_is_better(label, group)
        columns.append({
            "key": f"c{i}",
            "index": i,
            "label": label,
            "group": group,
            "category": cat,
            "role": role,
            "lower_is_better": lower,
            "target": PERSITE_TARGET if is_persite else None,
            "target_source": "persite_ratio" if is_persite else "target_row",
            "format": _col_format(label, cat, role),
        })

    if not columns:
        return None

    # ---- 目标值 ----
    # 两张表的写法不同，必须兼容：
    #   WOs SLA     目标行**按组横向合并**（B3:D3 = 0.98），只有组首列有值；
    #   Contract KPIs 目标行**每列各有一个值**（TTR 组 F..J = 0.98/0.98/0.95/0.9/0.9）。
    # 所以先按组跨度算一个兜底值，再让「本列自己的值」优先 —— 两种写法都得到正确结果。
    spans = _group_spans(groups)
    group_target: Dict[int, Optional[float]] = {}
    for _name, start, end in spans:
        target = None
        for i in range(start, min(end + 1, ncols)):
            target = _to_num(_cell(raw, target_row, i))
            if target is not None:
                break
        for i in range(start, min(end + 1, ncols)):
            group_target[i] = target
    for col in columns:
        if col["target_source"] == "persite_ratio":
            continue  # PerSite 用统一的 90% 占比目标，不读目标行
        own = _to_num(_cell(raw, target_row, col["index"]))
        col["target"] = own if own is not None else group_target.get(col["index"])

    # ---- 表头渲染行（服务端算好 colSpan，前端直接铺）----
    levels: List[Tuple[str, List[str]]] = []
    if cat_row is not None:
        levels.append(("category", cats))
    levels.append(("group", groups))
    levels.append(("column", leaves))
    header_rows: List[List[Dict[str, Any]]] = []
    for li, (lname, values) in enumerate(levels):
        is_leaf = li == len(levels) - 1
        cells: List[Dict[str, Any]] = []
        i = 1
        while i < ncols:
            if not is_leaf and not _to_text(values[i]) and not _to_text(leaves[i]):
                i += 1
                continue
            text = _to_text(values[i]) if not is_leaf else (leaves[i] or groups[i])
            span = 1
            if not is_leaf:
                # 同一父节点下同名连续列合并
                j = i + 1
                while (j < ncols and _to_text(values[j]) == _to_text(values[i])
                       and (li == 0 or _to_text(cats[j]) == _to_text(cats[i]))):
                    span += 1
                    j += 1
            cells.append({"text": text, "colSpan": span, "level": lname})
            i += span
        header_rows.append(cells)

    # ---- 数据行（Month 汇总行 + 日行）----
    data_start = int(cfg["data_start"])
    month_values: Dict[str, Any] = {}
    rows: List[Dict[str, Any]] = []
    for r in range(data_start, len(raw)):
        label_raw = _cell(raw, r, 0)
        label = _period_label(label_raw) if isinstance(label_raw, (datetime, date)) else _to_text(label_raw)
        if not label:
            continue
        values: Dict[str, Any] = {}
        blank = True
        for col in columns:
            v = _json_value(_cell(raw, r, col["index"]))
            values[col["key"]] = v
            if v is not None:
                blank = False
        if blank:
            continue
        if label.lower() in ("month", "total", "average", "avg"):
            month_values = values
            continue
        rows.append({"label": label, "values": values})

    # ---- 指标（metric）：叶子是角色列 -> 整组一个指标；否则每列一个指标 ----
    metrics: List[Dict[str, Any]] = []
    role_cols = [c for c in columns if c["role"] != "value"]
    if role_cols:
        by_group: Dict[str, List[Dict[str, Any]]] = {}
        order: List[str] = []
        for col in columns:
            gkey = f"{col['category']}|{col['group']}"
            if gkey not in by_group:
                by_group[gkey] = []
                order.append(gkey)
            by_group[gkey].append(col)
        for gkey in order:
            cols = by_group[gkey]
            value_col = next((c for c in cols if c["role"] == "achieve"), None) or cols[0]
            total_col = next((c for c in cols if c["role"] == "total"), None)
            within_col = next((c for c in cols if c["role"] == "within"), None)
            metrics.append(_metric_from(
                cols[0]["category"], cols[0]["group"], value_col, total_col, within_col,
                rows, month_values))
    else:
        for col in columns:
            metrics.append(_metric_from(
                col["category"], col["group"], col, None, None, rows, month_values))

    return {
        "sheet": cfg.get("sheet"),
        "label_header": label_header,
        "header_rows": header_rows,
        "columns": columns,
        "rows": rows,
        "month": month_values,
        "month_label": "Month",
        "metrics": metrics,
        "title": _to_text(_cell(raw, int(cfg["title_row"]), 0)) if cfg.get("title_row") is not None else "",
    }


def _metric_from(
    category: str,
    group: str,
    value_col: Dict[str, Any],
    total_col: Optional[Dict[str, Any]],
    within_col: Optional[Dict[str, Any]],
    rows: List[Dict[str, Any]],
    month_values: Dict[str, Any],
) -> Dict[str, Any]:
    """由取值列 + 明细行算出单个指标的统计（月值 / 达标天数 / min / max / avg）。"""
    key = value_col["key"]
    lower = bool(value_col["lower_is_better"])
    target = value_col["target"]

    series = [_to_num(r["values"].get(key)) for r in rows]
    present = [v for v in series if v is not None]
    met_flags = [_met(v, target, lower) for v in series]
    met_days = sum(1 for m in met_flags if m is True)
    below_days = sum(1 for m in met_flags if m is False)

    month_value = _to_num(month_values.get(key))
    label = group if value_col["role"] != "value" else value_col["label"]
    # 完整路径与矩阵列名/导出表头同口径（相邻同名只留一个），避免角色列只剩一个 "Achieve"
    parts = [p for p in (category, group, value_col["label"]) if p]
    full_parts: List[str] = []
    for p in parts:
        if not full_parts or full_parts[-1] != p:
            full_parts.append(p)

    return {
        "key": key,
        "label": label,
        "label_full": " / ".join(full_parts) or label,
        "category": category,
        "group": group,
        "target": target,
        "target_source": value_col.get("target_source", "target_row"),
        "format": value_col.get("format", "number"),
        "lower_is_better": lower,
        "value_key": key,
        "total_key": total_col["key"] if total_col else None,
        "within_key": within_col["key"] if within_col else None,
        "month_value": _round(month_value),
        "month_total": _round(_to_num(month_values.get(total_col["key"]))) if total_col else None,
        "month_within": _round(_to_num(month_values.get(within_col["key"]))) if within_col else None,
        "met": _met(month_value, target, lower),
        "day_count": len(present),
        "met_days": met_days,
        "below_days": below_days,
        "min": _round(min(present)) if present else None,
        "max": _round(max(present)) if present else None,
        "avg": _round(sum(present) / len(present)) if present else None,
    }


# =============================================================================
# 解析：单行表头的扁平表（FineBI Dashboard Planing / 利比里亚自动化报表）
# =============================================================================
def _parse_flat(raw: List[List[Any]], cfg: Dict[str, Any], sheet: str) -> Optional[Dict[str, Any]]:
    header_row = int(cfg["header_row"])
    start_col = int(cfg.get("start_col", 0))
    if len(raw) <= header_row:
        return None

    header = [_to_text(c) for c in raw[header_row][start_col:]]
    while header and not header[-1]:
        header.pop()
    columns = _normalize_columns(header)
    if not columns:
        return None

    # 合并单元格只保留左上角值：纵向前向填充需要补的列
    fill_idx = {i for i, name in enumerate(columns) if name in set(cfg.get("vertical_fill", ()))}
    last_seen: Dict[int, Any] = {}

    rows: List[List[Any]] = []
    for r in range(int(cfg["data_start"]), len(raw)):
        cells = list(raw[r][start_col:start_col + len(columns)])
        if len(cells) < len(columns):
            cells.extend([None] * (len(columns) - len(cells)))
        # 先判空再纵向填充 —— 否则最后一行空行会被上一行的 Domain 填成「有值」
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue
        for i in fill_idx:
            if cells[i] is None or (isinstance(cells[i], str) and not cells[i].strip()):
                cells[i] = last_seen.get(i)
            else:
                last_seen[i] = cells[i]
        rows.append(cells)

    title = ""
    if cfg.get("title_row") is not None:
        title = _to_text(_cell(raw, int(cfg["title_row"]), int(cfg.get("title_col", 0))))

    return {
        "sheet": sheet,
        "layout": "flat",
        "columns": columns,
        "rows": rows,
        "title": title,
    }


# =============================================================================
# 组装：一次读取 -> 结构化
# =============================================================================
class _Data:
    __slots__ = ("kind", "source", "matrices", "flats", "order", "warnings")

    def __init__(self, kind: str, source: Dict[str, Any],
                 matrices: Dict[str, Dict[str, Any]], flats: Dict[str, Dict[str, Any]],
                 order: List[str], warnings: List[str]):
        self.kind = kind
        self.source = source
        self.matrices = matrices
        self.flats = flats
        self.order = order
        self.warnings = warnings


def _read(kind: str, path: Path) -> _Data:
    names = REPORT_KINDS[kind]["sheets"]
    raw_map = _read_sheets(path, names)
    warnings: List[str] = []

    matrices: Dict[str, Dict[str, Any]] = {}
    flats: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []
    for name in names:
        if name not in raw_map:
            warnings.append(f"工作簿缺少 sheet：{name}")
            continue
        cfg = dict(SHEET_LAYOUT.get(name) or {})
        cfg["sheet"] = name
        raw = raw_map[name]
        if (cfg.get("layout") or "flat") == "kpi_matrix":
            parsed = _parse_kpi_matrix(raw, cfg)
            if parsed is None:
                warnings.append(f"sheet {name} 解析为空")
                continue
            matrices[name] = parsed
        else:
            parsed = _parse_flat(raw, cfg, name)
            if parsed is None:
                warnings.append(f"sheet {name} 解析为空")
                continue
            flats[name] = parsed
        order.append(name)

    stat = path.stat()
    source = {
        "kind": kind,
        "kind_label_en": REPORT_KINDS[kind]["label_en"],
        "kind_label_zh": REPORT_KINDS[kind]["label_zh"],
        "task": REPORT_KINDS[kind]["task"],
        "file": path.name,
        "path": str(path),
        "generated_at": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "size": stat.st_size,
    }
    return _Data(kind, source, matrices, flats, order, warnings)


_CACHE_LOCK = threading.Lock()
#: kind -> {"key": cache_key, "data": _Data}
_CACHE: Dict[str, Dict[str, Any]] = {}


def _load(kind: str = DEFAULT_KIND, force: bool = False) -> _Data:
    kind = normalize_kind(kind)
    path = latest_file(kind)
    if path is None:
        raise FileNotFoundError(
            f"未找到 {REPORT_KINDS[kind]['label_en']} 工作簿"
            f"（{MS_SLA_DIR} 下匹配 '{REPORT_KINDS[kind]['pattern']}' 的文件）"
        )
    key = _cache_key([path])
    slot = _CACHE.get(kind)
    if not force and slot and slot.get("key") == key and slot.get("data") is not None:
        return slot["data"]
    data = _read(kind, path)
    with _CACHE_LOCK:
        _CACHE[kind] = {"key": key, "data": data}
    return data


def invalidate_cache(kind: Optional[str] = None) -> None:
    """清空解析缓存（测试 / 排障用；正常请求靠 mtime 自动失效）。"""
    with _CACHE_LOCK:
        if kind is None:
            _CACHE.clear()
        else:
            _CACHE.pop(normalize_kind(kind), None)


# =============================================================================
# 统计
# =============================================================================
def _matrix_overview(matrix: Dict[str, Any]) -> Dict[str, Any]:
    """单个矩阵的 KPI 达标概览。"""
    metrics = matrix["metrics"]
    met = [m for m in metrics if m["met"] is True]
    below = [m for m in metrics if m["met"] is False]
    unknown = [m for m in metrics if m["met"] is None]
    days = [r["label"] for r in matrix["rows"]]

    def gap(m: Dict[str, Any]) -> float:
        """距目标的相对差距（用于「最需要关注」排序）。"""
        v, t = m.get("month_value"), m.get("target")
        if v is None or t in (None, 0):
            return 0.0
        return (t - v) / abs(t) if not m["lower_is_better"] else (v - t) / abs(t)

    ranked = sorted((m for m in metrics if m["met"] is not None),
                    key=gap, reverse=True)
    return {
        "kpi_count": len(metrics),
        "met_count": len(met),
        "below_count": len(below),
        "unknown_count": len(unknown),
        "days": len(days),
        "first_day": days[0] if days else "",
        "last_day": days[-1] if days else "",
        "worst": [dict(m, gap=_round(gap(m), 4)) for m in ranked[:5] if gap(m) > 0],
        "best": [dict(m, gap=_round(gap(m), 4)) for m in ranked[::-1][:5] if gap(m) <= 0],
    }


def _sheet_meta(data: _Data, name: str) -> Dict[str, Any]:
    if name in data.matrices:
        m = data.matrices[name]
        return {
            "sheet": name,
            "key": name,
            "label_key": SHEET_LABEL_KEY.get(name, name),
            "layout": "kpi_matrix",
            "records": len(m["rows"]),
            "columns": len(m["columns"]),
            "metrics": len(m["metrics"]),
            "days": len(m["rows"]),
        }
    f = data.flats[name]
    return {
        "sheet": name,
        "key": name,
        "label_key": SHEET_LABEL_KEY.get(name, name),
        "layout": "flat",
        "records": len(f["rows"]),
        "columns": len(f["columns"]),
        "metrics": 0,
        "days": 0,
    }


def get_summary(kind: str = DEFAULT_KIND) -> Dict[str, Any]:
    """页面统计卡片 + KPI 矩阵 + 分表元信息 + 报表标题。"""
    data = _load(kind)
    sheets = [_sheet_meta(data, n) for n in data.order]

    overview: Dict[str, Any] = {}
    matrices: Dict[str, Any] = {}
    for name, m in data.matrices.items():
        matrices[name] = {
            "sheet": name,
            "label_key": SHEET_LABEL_KEY.get(name, name),
            "title": m["title"],
            "label_header": m["label_header"],
            "header_rows": m["header_rows"],
            "columns": m["columns"],
            "rows": m["rows"],
            "month": m["month"],
            "month_label": m["month_label"],
            "metrics": m["metrics"],
            "overview": _matrix_overview(m),
        }

    if data.kind == "wo_sla":
        m = matrices.get("WOs SLA achievement Report") or (next(iter(matrices.values())) if matrices else None)
        if m:
            overview = dict(m["overview"])
            overall = m["metrics"][0] if m["metrics"] else {}
            overview.update({
                "month_total": overall.get("month_total"),
                "month_within": overall.get("month_within"),
                "month_achieve": overall.get("month_value"),
                "metric_count": len(m["metrics"]),
            })
    else:
        primary = matrices.get("Contract KPIs Report (New)")
        wo = matrices.get("WOs SLA achievement Report")
        if primary:
            overview = dict(primary["overview"])
            cats: List[Dict[str, Any]] = []
            for cat in dict.fromkeys(m["category"] for m in primary["metrics"]):
                group = [m for m in primary["metrics"] if m["category"] == cat]
                cats.append({
                    "category": cat or "-",
                    "count": len(group),
                    "met_count": sum(1 for m in group if m["met"] is True),
                    "below_count": sum(1 for m in group if m["met"] is False),
                })
            overview["categories"] = cats
        if wo:
            overview["wo_sla"] = wo["overview"]
        overview["flat_counts"] = {n: len(data.flats[n]["rows"]) for n in data.flats}

    notes: List[str] = []
    for name in data.order:
        title = ""
        if name in data.matrices:
            title = data.matrices[name]["title"]
        elif name in data.flats:
            title = data.flats[name]["title"]
        if title:
            notes.append(f"{name}: {title}")

    return {
        "kind": data.kind,
        "kind_label_en": REPORT_KINDS[data.kind]["label_en"],
        "kind_label_zh": REPORT_KINDS[data.kind]["label_zh"],
        "source": data.source,
        "overview": overview,
        "matrices": matrices,
        "sheets": sheets,
        "notes": notes,
        "warnings": data.warnings,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


# =============================================================================
# 分页查询（矩阵 sheet 压平成「一行一天」，flat sheet 原样）
# =============================================================================
def _parse_filters(raw: Optional[str]) -> Dict[str, List[str]]:
    if not raw:
        return {}
    try:
        obj = json.loads(raw)
    except (ValueError, TypeError):
        logger.warning("[ms_kpi_report] filters 不是合法 JSON，已忽略: %r", str(raw)[:200])
        return {}
    if not isinstance(obj, dict):
        return {}
    out: Dict[str, List[str]] = {}
    for k, v in obj.items():
        if not isinstance(k, str):
            continue
        vals = [_to_text(x) for x in (v if isinstance(v, (list, tuple)) else [v])]
        vals = [x for x in vals if x]
        if vals:
            out[k] = vals
    return out


def _matrix_table(matrix: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    """矩阵 -> 扁平表（列名 = ``Category / Group / Label``，行 = 每天一行）。"""
    names: List[str] = []
    for col in matrix["columns"]:
        parts = [p for p in (col["category"], col["group"], col["label"]) if p]
        # 去掉相邻重复（角色列时 group 与 label 不同；值列时 group 可能与 label 同名）
        dedup: List[str] = []
        for p in parts:
            if not dedup or dedup[-1] != p:
                dedup.append(p)
        names.append(" / ".join(dedup))
    columns = _normalize_columns([matrix["label_header"]] + names)

    rows: List[List[Any]] = []
    if matrix["month"]:
        rows.append([matrix["month_label"]]
                    + [matrix["month"].get(c["key"]) for c in matrix["columns"]])
    for r in matrix["rows"]:
        rows.append([r["label"]] + [r["values"].get(c["key"]) for c in matrix["columns"]])
    return columns, rows


def _table_for(data: _Data, sheet: str) -> Optional[Tuple[List[str], List[List[Any]]]]:
    if sheet in data.matrices:
        return _matrix_table(data.matrices[sheet])
    if sheet in data.flats:
        f = data.flats[sheet]
        return list(f["columns"]), [list(r) for r in f["rows"]]
    return None


def _numeric_flags(columns: List[str], rows: List[List[Any]]) -> List[bool]:
    sample = rows[:200]
    flags: List[bool] = []
    for i in range(len(columns)):
        seen = False
        all_num = True
        for r in sample:
            v = r[i] if i < len(r) else None
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            seen = True
            if _to_num(v) is None:
                all_num = False
                break
        flags.append(bool(seen and all_num))
    return flags


def _sort_key(value: Any, numeric: bool):
    if numeric:
        n = _to_num(value)
        return (n is None, n if n is not None else 0.0)
    s = _to_text(value)
    return (s == "", s.lower())


def get_rows(
    kind: str = DEFAULT_KIND,
    sheet: str = "",
    page: int = 1,
    page_size: int = DEFAULT_PAGE_SIZE,
    q: str = "",
    sort: str = "",
    order: str = "asc",
    filters: Optional[str] = None,
) -> Dict[str, Any]:
    """某 sheet 的分页数据（服务端筛选 + 排序 + 分页）。

    矩阵 sheet 会被压平成「一行一天、列名 = Category / Group / Label」，
    便于页面提供搜索/排序/导出；卡片与矩阵视图仍走 ``/summary``。
    """
    data = _load(kind)
    table = _table_for(data, sheet)
    if table is None:
        return {
            "kind": data.kind, "sheet": sheet, "columns": [], "numeric_columns": [],
            "rows": [], "total": 0, "page": 1, "page_size": page_size,
            "notes": [], "warnings": [
                f"未知的 sheet：{sheet}（可用：{', '.join(data.order) or '无'}）"],
        }
    columns, all_rows = table
    flags = _numeric_flags(columns, all_rows)

    flt = _parse_filters(filters)
    idx = {c: i for i, c in enumerate(columns)}
    active = [(idx[c], set(vals)) for c, vals in flt.items() if c in idx]
    needle = q.strip().lower()
    # 矩阵 sheet 压平后是「一行一天、列名才是 KPI 名」，只搜单元格值等于搜不到任何 KPI
    # （值全是数字/日期）。所以矩阵 sheet 额外允许命中列名 —— 搜 "TTR" 返回它的整条日序列。
    header_hit = (bool(needle) and sheet in data.matrices
                  and any(needle in c.lower() for c in columns))

    matched: List[List[Any]] = []
    for r in all_rows:
        ok = True
        for i, allowed in active:
            if _to_text(r[i] if i < len(r) else None) not in allowed:
                ok = False
                break
        if not ok:
            continue
        if needle and not header_hit and not any(
                needle in _to_text(v).lower() for v in r if v is not None):
            continue
        matched.append(r)

    if sort and sort in idx:
        i = idx[sort]
        numeric = flags[i]
        matched = sorted(
            matched,
            key=lambda r: _sort_key(r[i] if i < len(r) else None, numeric),
            reverse=order.lower() == "desc",
        )

    page_size = max(1, min(int(page_size or DEFAULT_PAGE_SIZE), MAX_PAGE_SIZE))
    total = len(matched)
    pages = max(1, (total + page_size - 1) // page_size)
    page = max(1, min(int(page or 1), pages))
    start = (page - 1) * page_size
    chunk = matched[start:start + page_size]

    rows = [
        {columns[i]: _json_value(r[i] if i < len(r) else None) for i in range(len(columns))}
        for r in chunk
    ]
    return {
        "kind": data.kind,
        "sheet": sheet,
        "columns": columns,
        "numeric_columns": [columns[i] for i, f in enumerate(flags) if f],
        "rows": rows,
        "total": total,
        "page": page,
        "page_size": page_size,
        "notes": [],
        "warnings": [],
    }


def get_filters(kind: str = DEFAULT_KIND, sheet: str = "") -> Dict[str, Any]:
    """某 sheet 的筛选下拉候选值（仅 flat sheet 有配置）。"""
    data = _load(kind)
    if sheet not in data.flats:
        return {"kind": data.kind, "sheet": sheet, "fields": [], "warnings": []}
    f = data.flats[sheet]
    idx = {c: i for i, c in enumerate(f["columns"])}
    fields: List[Dict[str, Any]] = []
    for column, label in FILTER_COLUMNS.get(sheet, ()):
        i = idx.get(column)
        if i is None:
            continue
        counter: Dict[str, int] = {}
        for r in f["rows"]:
            v = _to_text(r[i] if i < len(r) else None)
            if v:
                counter[v] = counter.get(v, 0) + 1
        options = [{"value": k, "count": v}
                   for k, v in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))[:60]]
        if options:
            fields.append({"column": column, "label": label, "options": options})
    return {"kind": data.kind, "sheet": sheet, "fields": fields, "warnings": []}


def get_sources(kind: str = DEFAULT_KIND) -> List[Dict[str, Any]]:
    data = _load(kind)
    out = [dict(data.source)]
    for name in data.order:
        out.append(_sheet_meta(data, name))
    return out
