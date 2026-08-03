"""导入 MW Topology Excel 到 SQLite"""
import sqlite3
import openpyxl
import time
import os
from backend.database import get_raw_connection

start = time.time()

DB_PATH = os.path.join(os.path.dirname(__file__), '..', 'db', 'app.db')
XLSX_PATH = os.path.join(os.path.dirname(__file__), '..', 'db', 'ZTE OLR MS MW Topology _OSS.xlsx')

with get_raw_connection(DB_PATH) as conn:

    # 创建 topology 表
    conn.execute("DROP TABLE IF EXISTS topology")
    conn.execute("""
    CREATE TABLE topology (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        site_a_id TEXT,
        site_a_name TEXT,
        site_b_id TEXT,
        site_b_name TEXT,
        hop_count INTEGER,
        hop_path TEXT,
        a_is_hub INTEGER DEFAULT 0,
        b_is_hub INTEGER DEFAULT 0
    )
""")
    conn.execute("CREATE INDEX idx_topology_site_a ON topology(site_a_id)")
    conn.execute("CREATE INDEX idx_topology_site_b ON topology(site_b_id)")

    # 获取 hub 站点集合
    hub_sites = set()
    for row in conn.execute("SELECT site_id FROM master_site_db WHERE hub_site = 'Y'"):
        hub_sites.add(str(row[0]))

    # 获取站点经纬度
    site_coords = {}
    for row in conn.execute("SELECT site_id, latitude, longitude FROM master_site_db WHERE latitude IS NOT NULL AND longitude IS NOT NULL"):
        site_coords[str(row[0])] = (row[1], row[2])

    # 读取 Excel
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    ws = wb.active

    imported = 0
    skipped = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        vals = [c.value for c in row]
        site_a_id = str(vals[0]) if vals[0] is not None else None
        site_a_name = vals[1] or ''
        site_b_id = str(vals[2]) if vals[2] is not None else None
        site_b_name = vals[3] or ''

        if not site_a_id or not site_b_id:
            skipped += 1
            continue

        # 收集 hop 路径
        hops = [str(v) for v in vals[4:25] if v is not None]
        hop_count = len(hops)
        hop_path = ','.join(hops)

        a_is_hub = 1 if site_a_id in hub_sites else 0
        b_is_hub = 1 if site_b_id in hub_sites else 0

        conn.execute(
            "INSERT INTO topology (site_a_id, site_a_name, site_b_id, site_b_name, hop_count, hop_path, a_is_hub, b_is_hub) VALUES (?,?,?,?,?,?,?,?)",
            (site_a_id, site_a_name, site_b_id, site_b_name, hop_count, hop_path, a_is_hub, b_is_hub)
        )
        imported += 1

    conn.commit()

    # 统计
    total = conn.execute("SELECT COUNT(*) FROM topology").fetchone()[0]
    hub_links = conn.execute("SELECT COUNT(*) FROM topology WHERE a_is_hub = 1 AND b_is_hub = 1").fetchone()[0]
    with_coords = conn.execute("""
    SELECT COUNT(*) FROM topology t
    WHERE t.site_a_id IN (SELECT site_id FROM master_site_db WHERE latitude IS NOT NULL)
    AND t.site_b_id IN (SELECT site_id FROM master_site_db WHERE latitude IS NOT NULL)
""").fetchone()[0]

    wb.close()
    conn.close()

    elapsed = time.time() - start
    print(f"Import done, elapsed: {elapsed:.2f}s")
    print(f"Total links: {total}, Hub-Hub links: {hub_links}, Links with coords: {with_coords}")
    print(f"Skipped: {skipped}")