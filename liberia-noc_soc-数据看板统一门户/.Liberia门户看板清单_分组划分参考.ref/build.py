# -*- coding: utf-8 -*-
"""从门户线上接口导出的目录 JSON 生成《Liberia 门户看板清单（分组划分参考）》。"""
import json, datetime

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


# ---- 商务蓝配色 ----
XL_HEADER_BG = xl_color("#4472C4")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_TITLE_FG = xl_color("#2F5597")
XL_INPUT_BG = xl_color("#D9E2F3")
XL_TOTAL_BG = xl_color("#2F5597")
XL_BORDER = xl_color("#BFBFBF")
XL_LINK = xl_color("#0563C1")
XL_OK_BG, XL_OK_FG = xl_color("#C6EFCE"), xl_color("#006100")
XL_WARN_BG, XL_WARN_FG = xl_color("#FFEB9C"), xl_color("#9C6500")

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SRC = r"D:\01_RPA\liberia-noc_soc-数据看板统一门户\TMP\live_portal.json"
OUT = r"D:\01_RPA\liberia-noc_soc-数据看板统一门户\Liberia门户看板清单_分组划分参考.xlsx"
TITLE = "Liberia NOC/SOC 门户看板清单（分组划分参考）"

PLATFORM_CN = {"custom": "自研 my_app", "finebi": "FineBI 报表", "external": "外部系统"}
STATUS_CN = {"online": "online", "beta": "beta"}

with open(SRC, encoding="utf-8") as f:
    portal = json.load(f)

groups = portal["groups"]
dashboards = portal["dashboards"]
# 按门户内分组顺序排列，组内保持原顺序
order = {g: i for i, g in enumerate(groups)}
dashboards = sorted(dashboards, key=lambda x: order.get(x.get("category"), 99))

# 修复源数据中的替换字符（原意为「交付达成率」）
def clean(s):
    return (s or "").replace("交付\ufffd\ufffd\ufffd成率", "交付达成率")

wb = Workbook()

# ============ Sheet 1: 看板清单 ============
ws = wb.active
ws.title = "看板清单"

HEADERS = [
    ("序号", 6), ("看板名称", 42), ("看板介绍", 62), ("当前归属模块", 16),
    ("平台类型", 13), ("URL 链接", 54), ("标签", 30), ("负责人", 14),
    ("归属部门", 18), ("状态", 9), ("更新频率", 12), ("访问量", 9),
    ("最近访问", 17), ("建议分组（待填写）", 18), ("分组调整理由 / 备注", 32),
]
HDR_ROW = 1
DATA_START = 2
DATA_END = DATA_START + len(dashboards) - 1

for c, (name, width) in enumerate(HEADERS, start=1):
    cell = ws.cell(row=HDR_ROW, column=c, value=name)
    cell.font = Font(bold=True, color=XL_HEADER_FG, size=10)
    cell.fill = PatternFill("solid", fgColor=XL_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(c)].width = width
ws.row_dimensions[HDR_ROW].height = 30

for i, d in enumerate(dashboards):
    r = DATA_START + i
    tags = "、".join(d.get("tags") or [])
    vals = [
        i + 1,
        clean(d.get("title")),
        clean(d.get("description")),
        clean(d.get("category")),
        PLATFORM_CN.get(d.get("platform"), d.get("platform") or ""),
        d.get("url") or "",
        tags,
        clean(d.get("owner")),
        clean(d.get("department")),
        STATUS_CN.get(d.get("status"), d.get("status") or ""),
        clean(d.get("updateFrequency")),
        d.get("clickCount") if isinstance(d.get("clickCount"), int) else None,
        (d.get("lastVisited") or "").replace("T", " "),
        "",  # 建议分组：留给评审人填写
        "",  # 备注：留给评审人填写
    ]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = BORDER
        cell.font = Font(size=10)
        if c in (3, 7):                            # 长文本
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        elif c in (1, 5, 10, 11, 12, 13):          # 短字段居中
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif c in (14, 15):                        # 输入区（评审人填写）
            cell.fill = PatternFill("solid", fgColor=XL_INPUT_BG)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    # 名称加粗；URL 超链接
    ws.cell(row=r, column=2).font = Font(size=10, bold=True)
    if vals[5]:
        u = ws.cell(row=r, column=6)
        u.hyperlink = vals[5]
        u.font = Font(size=9, color=XL_LINK, underline="single")
    ws.cell(row=r, column=12).number_format = "#,##0"

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A{HDR_ROW}:{get_column_letter(len(HEADERS))}{DATA_END}"

# 状态列条件格式（beta 试运行黄、online 在线绿）
status_range = f"J{DATA_START}:J{DATA_END}"
ws.conditional_formatting.add(status_range, CellIsRule(
    operator="equal", formula=['"beta"'], fill=PatternFill("solid", fgColor=XL_WARN_BG),
    font=Font(color=XL_WARN_FG, size=10)))
ws.conditional_formatting.add(status_range, CellIsRule(
    operator="equal", formula=['"online"'], fill=PatternFill("solid", fgColor=XL_OK_BG),
    font=Font(color=XL_OK_FG, size=10)))

# 建议分组下拉
dv = DataValidation(
    type="list",
    formula1='"' + ",".join(groups + ["新建分组"]) + '"',
    allow_blank=True, showDropDown=False)
dv.prompt = "从现有模块中选择，或选择「新建分组」"
dv.promptTitle = "建议分组"
ws.add_data_validation(dv)
dv.add(f"N{DATA_START}:N{DATA_END}")

# ============ Sheet 2: 分组汇总 ============
ws2 = wb.create_sheet("分组汇总")
S2_TITLE, S2_HDR = 1, 2
S2_START = 3
S2_END = S2_START + len(groups) - 1
S2_TOTAL = S2_END + 1

ws2.merge_cells(f"A{S2_TITLE}:H{S2_TITLE}")
t = ws2.cell(row=S2_TITLE, column=1, value="各模块看板数量与平台分布（现状）")
t.font = Font(bold=True, size=13, color=XL_TITLE_FG)
t.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[S2_TITLE].height = 26

S2_HEADERS = ["当前归属模块", "看板数量", "在线 online", "试运行 beta",
              "自研 my_app", "FineBI 报表", "外部系统", "占比"]
for c, name in enumerate(S2_HEADERS, start=1):
    cell = ws2.cell(row=S2_HDR, column=c, value=name)
    cell.font = Font(bold=True, color=XL_HEADER_FG, size=10)
    cell.fill = PatternFill("solid", fgColor=XL_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
ws2.row_dimensions[S2_HDR].height = 30

CAT = f"看板清单!$D${DATA_START}:$D${DATA_END}"
STA = f"看板清单!$J${DATA_START}:$J${DATA_END}"
PLT = f"看板清单!$E${DATA_START}:$E${DATA_END}"

for i, g in enumerate(groups):
    r = S2_START + i
    ws2.cell(row=r, column=1, value=g).alignment = Alignment(vertical="center", wrap_text=True)
    ws2.cell(row=r, column=2, value=f"=COUNTIF({CAT},$A{r})")
    ws2.cell(row=r, column=3, value=f'=COUNTIFS({CAT},$A{r},{STA},"online")')
    ws2.cell(row=r, column=4, value=f'=COUNTIFS({CAT},$A{r},{STA},"beta")')
    ws2.cell(row=r, column=5, value=f'=COUNTIFS({CAT},$A{r},{PLT},"自研 my_app")')
    ws2.cell(row=r, column=6, value=f'=COUNTIFS({CAT},$A{r},{PLT},"FineBI 报表")')
    ws2.cell(row=r, column=7, value=f'=COUNTIFS({CAT},$A{r},{PLT},"外部系统")')
    ws2.cell(row=r, column=8, value=f'=IF(OR($B${S2_TOTAL}="",$B${S2_TOTAL}=0),"",$B{r}/$B${S2_TOTAL})')
    for c in range(1, 9):
        cell = ws2.cell(row=r, column=c)
        cell.border = BORDER
        if c > 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "0.0%" if c == 8 else "#,##0"
        else:
            cell.font = Font(size=10)

ws2.cell(row=S2_TOTAL, column=1, value="合计")
for c in range(2, 8):
    col = get_column_letter(c)
    ws2.cell(row=S2_TOTAL, column=c, value=f"=SUM({col}{S2_START}:{col}{S2_END})")
ws2.cell(row=S2_TOTAL, column=8,
         value=f'=IF(OR($B${S2_TOTAL}="",$B${S2_TOTAL}=0),"",$B${S2_TOTAL}/$B${S2_TOTAL})')
for c in range(1, 9):
    cell = ws2.cell(row=S2_TOTAL, column=c)
    cell.font = Font(bold=True, color=XL_HEADER_FG, size=10)
    cell.fill = PatternFill("solid", fgColor=XL_TOTAL_BG)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if c == 8:
        cell.number_format = "0.0%"
    elif c > 1:
        cell.number_format = "#,##0"

for c, w in zip("ABCDEFGH", [22, 11, 12, 12, 13, 13, 11, 10]):
    ws2.column_dimensions[c].width = w
ws2.freeze_panes = f"A{S2_START}"

chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "各模块看板数量"
chart.y_axis.title = "看板数量"
chart.x_axis.title = None
chart.add_data(Reference(ws2, min_col=2, min_row=S2_HDR, max_row=S2_END), titles_from_data=True)
chart.set_categories(Reference(ws2, min_col=1, min_row=S2_START, max_row=S2_END))
chart.height, chart.width = 9.5, 15
chart.legend = None
ws2.add_chart(chart, f"A{S2_TOTAL + 2}")

# ============ Sheet 3: 说明 ============
ws3 = wb.create_sheet("说明")
rows = [
    ("数据来源", "门户线上目录接口 http://localhost/uniplam/api/portal-default（服务端 portal-default.json）"),
    ("门户入口", "http://localhost/uniplam/  ·  局域网 http://192.168.100.20:9267/uniplam/"),
    ("抓取时间", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ("目录保存时间", portal.get("savedAt", "")),
    ("看板总数", f"{len(dashboards)} 个"),
    ("当前分组数", f"{len(groups)} 个：" + "、".join(groups)),
    ("", ""),
    ("【字段说明】", ""),
    ("序号", "按门户当前分组顺序排列的流水号"),
    ("看板名称", "门户卡片标题，括号内为中文释义"),
    ("看板介绍", "门户卡片描述，说明看板口径、数据源与主要能力"),
    ("当前归属模块", "门户中该看板当前所属的 GROUP（分组）"),
    ("平台类型", "自研 my_app＝192.168.100.20/myapp 自研页面；FineBI 报表＝192.168.100.19 决策报表链接；外部系统＝其他独立系统"),
    ("URL 链接", "卡片实际跳转地址（自研页面链接含统一访问 token，可直接点击打开）"),
    ("标签", "门户卡片标签，可用于判断业务主题归属"),
    ("负责人 / 归属部门", "门户中登记的看板负责人与责任部门"),
    ("状态", "online＝正式在线；beta＝试运行/链接待更新（FineBI 侧链接多为待更新）"),
    ("更新频率", "看板数据刷新或报表更新节奏"),
    ("访问量 / 最近访问", "门户统计的累计点击次数与最近一次访问时间，0 或空表示暂无访问记录"),
    ("建议分组（待填写）", "评审人填写建议归入的 GROUP，单元格已内置下拉可选现有模块或「新建分组」"),
    ("分组调整理由 / 备注", "评审人填写调整理由、拆分/合并建议等"),
    ("", ""),
    ("【划分 GROUP 时可参考的信号】", ""),
    ("1. 业务域 vs 技术域", "告警/故障单属事件与工单流程，网络性能/传输承载属性能与资源域，MS 运维管理为考核与总览域，需先确定是否按此分层"),
    ("2. 使用角色", "一线值班（告警、故障单、GIS）与二线专家（KPI、核心网、传输）与高管（CXO、Overview）关注面差异明显"),
    ("3. 数据来源与平台", "FineBI 报表多为 beta 且链接待更新，自研 my_app 为正式在线，可考虑按「可用性」或「平台」再分层"),
    ("4. 通用工具类", "自动化调度与数据集成、其他 两组多为平台级工具（任务管理、表浏览器、帮助、反馈），通常不宜与业务看板混排"),
]
ws3.column_dimensions["A"].width = 26
ws3.column_dimensions["B"].width = 110
for i, (k, v) in enumerate(rows, start=1):
    a = ws3.cell(row=i, column=1, value=k)
    b = ws3.cell(row=i, column=2, value=v)
    b.alignment = Alignment(vertical="top", wrap_text=True)
    if k.startswith("【"):
        a.font = Font(bold=True, size=11, color=XL_TITLE_FG)
        b.font = Font(bold=True, size=11, color=XL_TITLE_FG)
    else:
        a.font = Font(bold=True, size=10)
        b.font = Font(size=10)

wb.properties.title = TITLE
wb.save(OUT)
print("OK ->", OUT)
print("dashboards:", len(dashboards), "groups:", len(groups))
print(f"看板清单 表头行={HDR_ROW} 数据行={DATA_START}..{DATA_END}")
print(f"分组汇总 标题行={S2_TITLE} 表头行={S2_HDR} 数据行={S2_START}..{S2_END} 合计行={S2_TOTAL}")
