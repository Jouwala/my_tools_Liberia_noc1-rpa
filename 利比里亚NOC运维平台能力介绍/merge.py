# -*- coding: utf-8 -*-
"""把 17 个单页 pptx 合并成一份完整演示文稿。

背景：本机 slidep 的保存服务当前对每个文件只允许成功落一页，
因此改为「每页单独生成一个 pptx」+「离线合并」的两步方案。
"""
import copy
import os
import sys

from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.opc.package import Part

R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
BASE = os.path.dirname(os.path.abspath(__file__))
media_cache = {}
media_seq = [0]


def copy_slide(src_path, dst):
    src = Presentation(src_path)
    s_slide = src.slides[0]

    try:
        layout = dst.slide_layouts[6]
    except Exception:
        layout = dst.slide_layouts[0]

    d_slide = dst.slides.add_slide(layout)
    # 清掉版式自带的占位符
    for ph in list(d_slide.shapes):
        ph._element.getparent().remove(ph._element)

    for shape in s_slide.shapes:
        d_slide.shapes._spTree.append(copy.deepcopy(shape._element))

    # 重映射引用关系（主要是图片）
    for el in d_slide.shapes._spTree.iter():
        for attr in ('embed', 'id', 'link'):
            rid = el.get(R_NS + attr)
            if not rid:
                continue
            try:
                rel = s_slide.part.rels[rid]
            except KeyError:
                continue
            if rel.is_external:
                new_rid = d_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel._target)
            else:
                tp = rel.target_part
                if not tp.content_type.startswith('image/'):
                    continue
                # slidep 会产出 svg+xml 分片，PIL 读不了，这里按原始字节复制
                key = (rel.reltype, tp.blob)
                if key in media_cache:
                    new_rid = media_cache[key]
                else:
                    ext = tp.content_type.split('/', 1)[1].replace('+xml', '')
                    media_seq[0] += 1
                    partname = PackURI('/ppt/media/m%03d.%s' % (media_seq[0], ext))
                    new_part = Part(partname, tp.content_type,
                                    d_slide.part.package, tp.blob)
                    new_rid = d_slide.part.relate_to(new_part, rel.reltype)
                    media_cache[key] = new_rid
            el.set(R_NS + attr, new_rid)
    return d_slide


def main():
    out = os.path.join(BASE, '利比里亚NOC运维平台能力介绍.pptx')
    pages = [os.path.join(BASE, 'p%02d.pptx' % i) for i in range(1, 18)]
    missing = [p for p in pages if not os.path.exists(p)]
    if missing:
        print('缺少单页文件：%s' % missing)
        return 1

    dst = Presentation(pages[0])
    for p in pages[1:]:
        copy_slide(p, dst)

    dst.save(out)
    print('已生成：%s（共 %d 页）' % (out, len(dst.slides.__iter__.__self__._sldIdLst)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
